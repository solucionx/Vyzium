import tempfile
import unittest
import re
from datetime import date, datetime, timedelta
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook

from engine import AutoScheduler, Store, WorkbookImporter, FollowUpService, classify, normalize, person_matches, phone_value


class ClassificationTests(unittest.TestCase):
    def test_due_soon_and_boundaries(self):
        today = date(2026, 9, 16)
        self.assertEqual(classify(today.isoformat(), None, "Solicitado", "", today=today).urgency, "due_soon")
        self.assertEqual(classify((today + timedelta(days=3)).isoformat(), None, "", "", today=today).urgency, "due_soon")
        self.assertEqual(classify((today + timedelta(days=4)).isoformat(), None, "", "", today=today).urgency, "scheduled")
        self.assertEqual(classify((today - timedelta(days=10)).isoformat(), None, "", "", today=today).urgency, "overdue")
        self.assertEqual(classify((today - timedelta(days=11)).isoformat(), None, "", "", today=today).urgency, "critical")

    def test_completed_and_cancelled_never_send(self):
        old = (date.today() - timedelta(days=50)).isoformat()
        self.assertFalse(classify(old, date.today().isoformat(), "", "").eligible)
        self.assertFalse(classify(old, None, "3 - Cancelado", "").eligible)

    def test_partial_receipt_remains_eligible(self):
        old = (date.today() - timedelta(days=2)).isoformat()
        result = classify(old, date.today().isoformat(), "1 - Recebido Parcialmente", "")
        self.assertTrue(result.eligible)
        self.assertEqual(result.urgency, "overdue")

    def test_partial_receipt_without_numeric_status_stays_open(self):
        old = (date.today() - timedelta(days=2)).isoformat()
        result = classify(
            old, date.today().isoformat(), "", "", order_status_code=None,
            received_qty=4, remaining_qty=6, quantity=10
        )
        self.assertTrue(result.eligible)
        self.assertEqual(result.urgency, "overdue")

    def test_numeric_status_is_authoritative(self):
        old = (date.today() - timedelta(days=2)).isoformat()
        requested = classify(old, date.today().isoformat(), "0 - Solicitado", "", order_status_code=0)
        completed = classify(old, None, "", "", order_status_code=2)
        self.assertTrue(requested.eligible)
        self.assertFalse(completed.eligible)

    def test_scheduler_boundary(self):
        settings = {"automatic_enabled": True, "simulation": False, "last_auto_run": "", "schedule_time": "09:00"}
        self.assertFalse(AutoScheduler.should_run(settings, __import__("datetime").datetime(2026, 9, 16, 8, 59)))
        self.assertTrue(AutoScheduler.should_run(settings, __import__("datetime").datetime(2026, 9, 16, 9, 0)))
        settings["last_auto_run"] = "2026-09-16"
        self.assertFalse(AutoScheduler.should_run(settings, __import__("datetime").datetime(2026, 9, 16, 12, 0)))

    def test_scheduler_retries_failed_day_with_one_hour_backoff(self):
        temp = tempfile.TemporaryDirectory()
        try:
            store = Store(Path(temp.name) / "scheduler.db")
            store.save_settings({"simulation": False, "automatic_enabled": True, "schedule_time": "09:00"})
            store.set_internal("last_workbook_path", str(Path(temp.name) / "base.xlsx"))

            class Importer:
                def import_file(self, _path):
                    return {"ok": True}

            class Service:
                def send(self, **_kwargs):
                    return {"sent": 0, "simulated": 0, "failed": 1, "uncertain": 0, "results": []}

            scheduler = AutoScheduler(store, Importer(), Service())
            now = datetime(2026, 9, 16, 9, 0)
            self.assertTrue(scheduler.run_once_if_due(now))
            settings = store.settings()
            self.assertNotEqual(settings.get("last_auto_run"), "2026-09-16")
            self.assertFalse(AutoScheduler.should_run(settings, now + timedelta(minutes=59)))
            self.assertTrue(AutoScheduler.should_run(settings, now + timedelta(minutes=61)))
            store.connection.close()
        finally:
            temp.cleanup()



class ImportTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.store = Store(self.root / "test.db")

    def tearDown(self):
        self.store.connection.close()
        self.temp.cleanup()

    def make_workbook(self):
        path = self.root / "base.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Base"
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "QUANTIDADEOC", "STATUSITEMDAORDEMDECOMPRA"])
        sheet.append([123, "Hotel A", "Filtro", date.today(), "Fornecedor Ágil", 2, "0 - Solicitado"])
        sheet.append([123, "Hotel A", "Refil", date.today(), "Fornecedor Ágil", 1, "0 - Solicitado"])
        workbook.save(path)
        return path

    def test_import_group_and_simulate(self):
        result = WorkbookImporter(self.store).import_file(str(self.make_workbook()))
        self.assertEqual(result["read"], 2)
        self.store.save_settings({"buyer_filter": ""})
        supplier = self.store.suppliers()[0]
        self.store.save_supplier({**supplier, "phone": "85999999999"})
        groups = FollowUpService(self.store).groups()
        self.assertEqual(len(groups), 1)
        self.assertEqual(groups[0]["order_count"], 1)
        self.assertEqual(groups[0]["item_count"], 2)
        sent = FollowUpService(self.store).send(force_simulation=True)
        self.assertEqual(sent["simulated"], 1)
        self.assertEqual(len(FollowUpService(self.store).groups()), 1)

    def test_normalization(self):
        self.assertEqual(normalize("Razão Social - LTDA"), "RAZAO SOCIAL LTDA")
        self.assertTrue(person_matches("COMPRADOR", "COMPRADOR TESTE"))
        self.assertTrue(person_matches("COMPRADOR TESTE", "COMPRADOR"))
        self.assertFalse(person_matches("MARIANA PEIXOTO", "COMPRADOR TESTE"))

    def test_brazilian_legacy_mobile_gets_ninth_digit(self):
        self.assertEqual(phone_value("+558599216923"), "+5585999216923")
        self.assertEqual(phone_value("85 9921-6923"), "+5585999216923")
        self.assertEqual(phone_value("85 3232-1234"), "+558532321234")

    def test_import_repairs_incorrect_excel_dimension(self):
        original = self.make_workbook()
        broken = self.root / "base-dimension-a1.xlsx"
        with ZipFile(original, "r") as source, ZipFile(broken, "w", ZIP_DEFLATED) as target:
            for entry in source.infolist():
                payload = source.read(entry.filename)
                if entry.filename == "xl/worksheets/sheet1.xml":
                    payload = re.sub(rb'<dimension ref="[^"]+"\s*/>', b'<dimension ref="A1"/>', payload)
                target.writestr(entry, payload)
        result = WorkbookImporter(self.store).import_file(str(broken))
        self.assertEqual(result["read"], 2)
        self.assertEqual(result["order_items"], 2)

    def test_invalid_workbook_has_friendly_error_and_preserves_data(self):
        WorkbookImporter(self.store).import_file(str(self.make_workbook()))
        invalid = self.root / "arquivo-errado.xlsx"
        workbook = Workbook()
        workbook.active.append(["COLUNA A", "COLUNA B"])
        workbook.active.append([1, 2])
        workbook.save(invalid)
        with self.assertRaisesRegex(ValueError, "BASE SCI.xlsx"):
            WorkbookImporter(self.store).import_file(str(invalid))
        self.assertEqual(self.store.current_order_count(), 2)

    def test_new_import_replaces_previous_snapshot(self):
        importer = WorkbookImporter(self.store)
        path = self.make_workbook()
        importer.import_file(str(path))
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR"])
        sheet.append([999, "Hotel B", "Item novo", date.today(), "Fornecedor Novo"])
        workbook.save(path)
        importer.import_file(str(path))
        orders = self.store.order_rows()
        self.assertEqual(len(orders), 1)
        self.assertEqual(orders[0]["oc"], "999")

    def test_manual_control_survives_daily_import(self):
        importer = WorkbookImporter(self.store)
        path = self.make_workbook()
        importer.import_file(str(path))
        supplier_key = self.store.order_rows()[0]["supplier_key"]
        saved = self.store.save_order_control({
            "oc": "123", "supplier_key": supplier_key, "control_status": "card_payment",
            "note": "Pagamento será feito no cartão corporativo."
        })
        self.assertEqual(saved["control_label"], "Pedido aguardando pagamento no cartão")

        importer.import_file(str(path))
        control = self.store.order_control("123", supplier_key)
        self.assertEqual(control["control_status"], "card_payment")
        self.assertEqual(control["note"], "Pagamento será feito no cartão corporativo.")
        self.assertEqual(len(control["history"]), 1)

        service = FollowUpService(self.store)
        by_control = service.order_summaries(control_status="card_payment")
        by_note = service.order_summaries(search="cartão corporativo")
        self.assertEqual(len(by_control), 1)
        self.assertEqual(len(by_note), 1)

    def test_marking_order_as_sent_records_date(self):
        WorkbookImporter(self.store).import_file(str(self.make_workbook()))
        supplier_key = self.store.order_rows()[0]["supplier_key"]
        saved = self.store.save_order_control({
            "oc": "123", "supplier_key": supplier_key, "control_status": "sent", "note": "Pedido enviado por e-mail."
        })
        self.assertIsNotNone(saved["sent_at"])
        self.assertEqual(saved["control_label"], "Pedido enviado")

    def test_base_sci_receipts_and_order_detail(self):
        path = self.root / "base-sci.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Query"
        sheet.append([
            "IDORDEMDECOMPRA", "FKEMPRESA", "EMPRESA", "IDSCI", "IDITEMDASCI", "CODIGOARTIGO",
            "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "FKFORNECEDOR",
            "QUANTIDADEOC", "STATUSITEMDAORDEMDECOMPRA", "NMSTATUSITEMDAORDEMDECOMPRA",
            "STATUSBPMOC", "NMSTATUSBPMOC", "QUANTIDADERECEBIDA", "DATAENTRADAMERCADORIA",
            "ANNUMERODANOTAFISCAL", "UNIDADEMEDIDARECEBIDA", "COMPRADOR", "URGENTE"
        ])
        common = [321, 7, "Hotel A", 100, 9001, "ABC", "Filtro", date.today(), "Fornecedor A", 55,
                  10, "1 - Recebido Parcialmente", 1, "3 - Aprovado", 3]
        sheet.append(common + [4, date.today(), "NF1", "UN", "COMPRADOR TESTE", "Não"])
        sheet.append(common + [6, date.today(), "NF2", "UN", "COMPRADOR TESTE", "Sim"])
        workbook.save(path)

        result = WorkbookImporter(self.store).import_file(str(path))
        self.assertEqual(result["order_items"], 1)
        self.assertEqual(result["receipts"], 2)
        item = self.store.order_rows()[0]
        self.assertEqual(item["received_qty"], 10)
        self.assertEqual(item["remaining_qty"], 0)
        self.assertEqual(item["urgent"], 1)

        service = FollowUpService(self.store)
        summary = service.order_summaries(buyer="COMPRADOR TESTE", urgency="open")[0]
        self.assertEqual(summary["warning_count"], 1)
        self.assertEqual(summary["attendance_status"], "partial")
        self.assertEqual(summary["attendance_label"], "Atendida parcialmente")
        self.assertEqual(len(service.order_summaries(attendance_status="partial")), 1)
        self.assertEqual(len(service.order_summaries(attendance_status=["pending", "partial"])), 1)
        self.assertEqual(service.order_summaries(attendance_status="attended"), [])
        detail = service.order_detail("321", summary["supplier_key"])
        self.assertEqual(len(detail["items"][0]["receipts"]), 2)

    def test_attendance_filter_accepts_multiple_statuses(self):
        path = self.root / "attendance-multi.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "QUANTIDADEOC", "STATUSITEMDAORDEMDECOMPRA"])
        sheet.append([1001, "Hotel A", "Item pendente", date.today(), "Fornecedor A", 1, "0 - Solicitado"])
        sheet.append([1002, "Hotel B", "Item parcial", date.today(), "Fornecedor B", 2, "1 - Recebido Parcialmente"])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        service = FollowUpService(self.store)
        self.assertEqual(len(service.order_summaries(attendance_status="pending")), 1)
        self.assertEqual(len(service.order_summaries(attendance_status="partial")), 1)
        combined = service.order_summaries(attendance_status=["pending", "partial"])
        self.assertEqual({row["attendance_status"] for row in combined}, {"pending", "partial"})
        self.assertEqual(len(combined), 2)

    def test_attendance_summary_states(self):
        service = FollowUpService(self.store)
        self.assertEqual(service._summary_attendance([{"order_status_code": 0}]), "pending")
        self.assertEqual(service._summary_attendance([{"order_status_code": 1}]), "partial")
        self.assertEqual(service._summary_attendance([{"order_status_code": 2}]), "attended")
        self.assertEqual(service._summary_attendance([{"order_status_code": 3}]), "canceled")
        self.assertEqual(service._summary_attendance([
            {"order_status_code": 2}, {"order_status_code": 0}
        ]), "partial")

    def test_supplier_phone_is_reused_for_every_order(self):
        path = self.root / "two-orders.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR"])
        sheet.append([100, "Hotel A", "Item A", date.today(), "Fornecedor Único"])
        sheet.append([101, "Hotel B", "Item B", date.today(), "Fornecedor Único"])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        supplier = self.store.suppliers()[0]
        self.store.save_supplier({**supplier, "phone": "85999999999", "contact_name": "Carlos"})
        self.store.save_supplier({"supplier_key": supplier["supplier_key"], "display_name": supplier["display_name"], "phone": "85888888888"})
        summaries = FollowUpService(self.store).order_summaries()
        self.assertEqual(len(summaries), 2)
        self.assertTrue(all(row["phone"] == "+5585888888888" for row in summaries))
        self.assertEqual(self.store.suppliers()[0]["contact_name"], "Carlos")

    def test_document_supplier_key_is_preserved_when_saving_phone(self):
        path = self.root / "supplier-doc.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "CPFCNPJFORNECEDOR"])
        sheet.append([888, "Hotel C", "Item", date.today(), "Fornecedor Documento", "47.869.266/0001-00"])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        canonical = self.store.order_rows()[0]["supplier_key"]
        self.assertEqual(canonical, "47.869.266/0001-00")
        saved = self.store.save_supplier({
            "supplier_key": canonical, "display_name": "Fornecedor Documento", "phone": "85999999999"
        })
        self.assertEqual(saved["supplier_key"], canonical)
        self.assertEqual(len(self.store.suppliers()), 1)
        self.assertEqual(FollowUpService(self.store).order_summaries()[0]["phone"], "+5585999999999")

    def test_old_normalized_supplier_alias_is_redirected_to_canonical_key(self):
        path = self.root / "supplier-alias.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "CPFCNPJFORNECEDOR"])
        sheet.append([889, "Hotel C", "Item", date.today(), "Fornecedor Documento", "47.869.266/0001-00"])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        saved = self.store.save_supplier({
            "supplier_key": "47 869 266 0001 00", "display_name": "Fornecedor Documento", "phone": "85999999999"
        })
        self.assertEqual(saved["supplier_key"], "47.869.266/0001-00")
        self.assertEqual(len(self.store.suppliers()), 1)

    def test_startup_migrates_phone_from_old_normalized_alias(self):
        path = self.root / "supplier-migrate.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append(["OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "CPFCNPJFORNECEDOR"])
        sheet.append([891, "Hotel C", "Item", date.today(), "Fornecedor Documento", "47.869.266/0001-00"])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        now = datetime.now().isoformat(timespec="seconds")
        with self.store.lock:
            self.store.connection.execute(
                "INSERT INTO suppliers(supplier_key,display_name,phone,contact_name,active,updated_at) VALUES(?,?,?,?,?,?)",
                ("47 869 266 0001 00", "Fornecedor Documento", "+5585999999999", "Daniel", 1, now),
            )
            self.store.connection.commit()
        db_path = self.root / "test.db"
        self.store.connection.close()
        repaired = Store(db_path)
        rows = repaired.suppliers()
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["supplier_key"], "47.869.266/0001-00")
        self.assertEqual(rows[0]["phone"], "+5585999999999")
        self.assertEqual(rows[0]["contact_name"], "Daniel")
        repaired.connection.close()
        # Reopen for tearDown compatibility.
        self.store = Store(db_path)

    def test_rejected_text_without_numeric_bpm_code_is_not_sent(self):
        path = self.root / "rejected-text.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "OC", "EMPRESA", "DESCRICAOARTIGO", "DATAPREVISTAENTREGAOC",
            "RAZAOSOCIALFORNECEDOR", "STATUSITEMDAORDEMDECOMPRA", "STATUSBPMOC"
        ])
        sheet.append([890, "Hotel D", "Item", date.today(), "Fornecedor D", "0 - Solicitado", "RECUSADA"])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        self.store.save_settings({"buyer_filter": ""})
        self.assertEqual(FollowUpService(self.store).groups(include_blocked=True), [])

    def test_rejected_order_is_visible_but_not_eligible_for_followup(self):
        path = self.root / "rejected.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.append([
            "IDORDEMDECOMPRA", "EMPRESA", "IDITEMDASCI", "DESCRICAOARTIGO",
            "DATAPREVISTAENTREGAOC", "RAZAOSOCIALFORNECEDOR", "STATUSITEMDAORDEMDECOMPRA",
            "NMSTATUSITEMDAORDEMDECOMPRA", "NMSTATUSBPMOC"
        ])
        sheet.append([777, "Hotel B", 1, "Item", date.today(), "Fornecedor B", "0 - Solicitado", 0, 4])
        workbook.save(path)
        WorkbookImporter(self.store).import_file(str(path))
        service = FollowUpService(self.store)
        self.assertEqual(len(service.order_summaries(urgency="open")), 1)
        self.assertEqual(service.groups(include_blocked=True), [])


if __name__ == "__main__":
    unittest.main()
