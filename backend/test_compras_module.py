import copy
import tempfile
import unittest
from unittest.mock import patch
from pathlib import Path
from openpyxl import Workbook
import compras_engine as engine


def sample():
    return {'id': '3|1|100', 'company': 'Hotel A', 'sci': '1', 'article': 'A01',
            'description': 'Lâmpada', 'quantity': '10', 'unit': 'UN', 'buyer': 'Comprador A',
            'group': 'Elétrica', 'needed': '2026-09-20', 'issued': '2026-09-01', 'urgent': False,
            'status': 'pending', 'approval': 'approved'}


class PurchasesTest(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.s = engine.Store(self.tmp.name)
        self.s.put('items', sample()['id'], sample())
        self.m = self.s.create_map({'name': 'Cotação teste', 'ids': [sample()['id']]})
        self.m['suppliers'] = [{'id': 's1', 'name': 'Fornecedor A', 'phone': '5585999999999'},
                               {'id': 's2', 'name': 'Fornecedor B', 'phone': '5585988888888'}]
        self.m['quotes'] = {sample()['id']: {'s1': {'price': '10', 'discount': '10', 'kind': 'percent'},
                                            's2': {'price': '9.50', 'discount': '0', 'kind': 'amount'}}}
        self.m = self.s.save_map(self.m)['map']

    def test_discount_changes_winner(self):
        result = engine.evaluate(self.m)
        self.assertEqual(result['lines'][0]['chosen']['supplier_id'], 's1')
        self.assertEqual(result['net'], '90.00')
        self.assertEqual(result['saving'], '10.00')

    def test_amount_discount_is_line_total(self):
        self.m['quotes'][sample()['id']]['s1'].update(kind='amount', discount='2')
        r = engine.evaluate(self.m)
        self.assertEqual(r['lines'][0]['quotes'][0]['net'], '98.00')
        self.assertEqual(r['lines'][0]['chosen']['supplier_id'], 's2')

    def test_empty_quote_never_wins(self):
        self.m['quotes'][sample()['id']]['s1']['price'] = ''
        self.assertEqual(engine.evaluate(self.m)['lines'][0]['chosen']['supplier_id'], 's2')

    def test_zero_and_excessive_discount_rejected(self):
        for price, discount, kind in [('0', '0', 'amount'), ('10', '101', 'percent'), ('10', '101', 'amount'), ('-1', '0', 'amount'), ('NaN', '0', 'amount')]:
            with self.subTest(price=price, discount=discount, kind=kind):
                self.m['quotes'][sample()['id']]['s1'].update(price=price, discount=discount, kind=kind)
                with self.assertRaises(ValueError):
                    engine.evaluate(self.m)

    def test_tie_requires_choice_and_invalid_choice_does_not_override_price(self):
        self.m['quotes'][sample()['id']]['s2']['price'] = '9'
        self.assertEqual(engine.evaluate(self.m)['ties'], 1)
        self.assertEqual(engine.evaluate(self.m)['net'], '0')
        self.m['choices'] = {sample()['id']: 's2'}
        self.assertEqual(engine.evaluate(self.m)['net'], '90.00')
        self.m['quotes'][sample()['id']]['s2']['price'] = '11'
        self.assertEqual(engine.evaluate(self.m)['lines'][0]['chosen']['supplier_id'], 's1')

    def test_locale_decimal_and_rounding(self):
        self.m['quotes'][sample()['id']]['s1']['price'] = '1.234,567'
        d = self.s.save_map(self.m)
        self.assertEqual(d['map']['quotes'][sample()['id']]['s1']['price'], '1234.567')
        self.assertEqual(d['result']['lines'][0]['quotes'][0]['gross'], '12345.67')

    def test_reimport_preserves_map_and_detects_changes(self):
        changed = sample()
        changed['quantity'] = '20'
        self.s.put('items', changed['id'], changed)
        d = self.s.detail(self.m['id'])
        self.assertEqual(d['map']['items'][0]['quantity'], '10')
        self.assertEqual(len(d['changes']), 1)
        with self.assertRaises(ValueError):
            self.s.preview(self.m['id'], 's1')

    def test_prevent_duplicate_maps(self):
        with self.assertRaises(ValueError):
            self.s.create_map({'name': 'Outro', 'ids': [sample()['id']]})

    def test_quantity_tampering_does_not_change_source(self):
        self.m['items'][0]['quantity'] = '999'
        self.assertEqual(self.s.save_map(self.m)['map']['items'][0]['quantity'], '10')

    def test_revision_prevents_lost_updates(self):
        old = copy.deepcopy(self.m)
        self.s.save_map(self.m)
        with self.assertRaises(ValueError):
            self.s.save_map(old)

    def send_body(self):
        p = self.s.preview(self.m['id'], 's1')
        return {'map_id': self.m['id'], 'supplier_id': 's1', 'fingerprint': p['fingerprint'], 'revision': p['revision']}

    def test_message_has_scope_without_competitor_prices(self):
        p = self.s.preview(self.m['id'], 's1')
        self.assertIn('Hotel: Hotel A | SCI: 1', p['message'])
        self.assertIn('Quantidade: 10 UN', p['message'])
        self.assertNotIn('Fornecedor B', p['message'])
        self.assertNotIn('90.00', p['message'])

    def test_sending_is_explicit_and_deduplicated(self):
        with patch('compras_engine.whatsapp_request', side_effect=[{'ready': True}, {'status': 'sent'}]) as bridge:
            self.assertEqual(self.s.send(self.send_body())['status'], 'sent')
            with self.assertRaises(ValueError):
                self.s.send(self.send_body())
            self.assertEqual(bridge.call_count, 2)

    def test_uncertain_result_requires_review(self):
        with patch('compras_engine.whatsapp_request', side_effect=[{'ready': True}, TimeoutError('timeout')]):
            msg = self.s.send(self.send_body())
        self.assertEqual(msg['status'], 'uncertain')
        with self.assertRaises(ValueError):
            self.s.send(self.send_body())
        self.s.review({'id': msg['id'], 'outcome': 'not_received'})
        with patch('compras_engine.whatsapp_request', side_effect=[{'ready': True}, {'status': 'sent'}]):
            self.assertEqual(self.s.send(self.send_body())['status'], 'sent')

    def test_connection_failure_is_not_uncertain(self):
        with patch('compras_engine.whatsapp_request', side_effect=RuntimeError('não conectado')):
            self.assertEqual(self.s.send(self.send_body())['status'], 'failed')

    def test_stale_preview_does_not_send(self):
        body = self.send_body()
        self.m['items'][0]['note'] = 'Nova especificação'
        self.s.save_map(self.m)
        with patch('compras_engine.whatsapp_request') as bridge, self.assertRaises(ValueError):
            self.s.send(body)
        bridge.assert_not_called()

    def test_csv_formula_injection_and_values(self):
        self.m['items'][0]['note'] = '=HYPERLINK("malicious")'
        self.s.save_map(self.m)
        csv = self.s.export_csv(self.m['id'])['content']
        self.assertIn("'=HYPERLINK", csv)
        self.assertIn('90,00', csv)

    def test_restart_recovers_interrupted_send(self):
        msg = {'id': 'interrupt', 'status': 'sending'}
        self.s.put('messages', msg['id'], msg)
        reopened = engine.Store(self.tmp.name)
        self.assertEqual(reopened.all('messages')[0]['status'], 'uncertain')


class ImportHardeningTest(unittest.TestCase):
    def test_header_can_start_after_report_metadata_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / 'base.xlsx'
            wb = Workbook()
            ws = wb.active
            ws.append(['RELATÓRIO DE SCI'])
            ws.append(['Gerado em', '20/09/2026'])
            ws.append(['FKEMPRESA', 'EMPRESA', 'IDSCI', 'IDITEMDASCI', 'DESCRICAOARTIGO', 'QUANTIDADESCI', 'UNIDADEDEMEDIDASCI', 'IDORDEMDECOMPRA', 'NMSTATUSDOITEMDASCI', 'NMSTATUSBPMSCI', 'COMPRADOR'])
            ws.append(['3', 'Hotel A', '1', '10', 'Item', 2, 'UN', '', '0', '3', 'Comprador A'])
            wb.save(path)
            items, report = engine.load_items(path)
            self.assertEqual(len(items), 1)
            self.assertEqual(report['eligible'], 1)

    def test_internal_json_table_name_is_allowlisted(self):
        with tempfile.TemporaryDirectory() as tmp:
            store = engine.Store(tmp)
            with self.assertRaisesRegex(ValueError, 'Tabela interna inválida'):
                store.all('items; DROP TABLE maps;--')
            with self.assertRaisesRegex(ValueError, 'Tabela interna inválida'):
                store.put('sqlite_master', 'x', {})


class ImportTest(unittest.TestCase):
    def test_import_deduplicates_and_blocks_ordered_closed_and_rejected_items(self):
        fields = ['FKEMPRESA', 'EMPRESA', 'IDSCI', 'IDITEMDASCI', 'DESCRICAOARTIGO', 'QUANTIDADESCI', 'UNIDADEDEMEDIDASCI', 'IDORDEMDECOMPRA', 'NMSTATUSDOITEMDASCI', 'NMSTATUSBPMSCI', 'COMPRADOR']
        rows = [['3', 'Hotel A', '1', '1', 'Item', 10, 'UN', '', '0', '3', 'Comprador A']]
        rows += [rows[0].copy()]
        for ident, oc, status, approval in [('2', '22', '0', '3'), ('2', '', '0', '3'), ('3', '', '6', '3'), ('4', '', '0', '2'), ('5', '', '2', '3'), ('6', '', '5', '3')]:
            rows.append(['3', 'Hotel A', '1', ident, 'Item', 10, 'UN', oc, status, approval, 'Comprador A'])
        class Sheet:
            def reset_dimensions(self): pass
            def iter_rows(self, **kwargs): return iter([fields] + rows)
        class Book:
            def __iter__(self): return iter([Sheet()])
            def close(self): pass
        with patch('compras_engine.load_workbook', return_value=Book()):
            items, report = engine.load_items('mock.xlsx')
        self.assertEqual(len(items), 2)
        self.assertEqual({i['id'] for i in items}, {'3|1|1', '3|1|5'})
        self.assertEqual(report['excluded']['with_order'], 1)
        self.assertEqual(report['excluded']['closed_or_rejected'], 2)
        self.assertEqual(report['excluded']['other_status'], 1)


if __name__ == '__main__':
    unittest.main()
