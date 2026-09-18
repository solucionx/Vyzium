from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sqlite3
import sys
import threading
import time
import unicodedata
import uuid
from dataclasses import dataclass
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from typing import Any
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel
from urllib.parse import parse_qs, urlparse
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError


def whatsapp_request(route, body=None):
    base = os.environ.get("FOLLOWUP_WHATSAPP_URL", "")
    if not re.fullmatch(r"http://127\.0\.0\.1:\d+", base):
        raise RuntimeError("Inicie o aplicativo Vyzium para usar a conexão do WhatsApp.")
    request = Request(base + route, data=json.dumps(body or {}).encode(), headers={
        "Content-Type": "application/json", "X-FollowUp-Token": os.environ.get("FOLLOWUP_API_TOKEN", "")
    }, method="POST")
    try:
        timeout = 130 if route == "/wait" else (10 if route == "/health" else 75)
        with urlopen(request, timeout=timeout) as response:
            return json.load(response)
    except HTTPError as exc:
        try:
            message = json.load(exc).get("error", "Erro na conexão do WhatsApp.")
        except Exception:
            message = "Erro na conexão do WhatsApp."
        raise RuntimeError(message) from exc
    except (URLError, TimeoutError, ConnectionError, json.JSONDecodeError) as exc:
        # Ponte fechada, processo do WhatsApp caiu ou resposta corrompida: nunca deixar
        # o erro técnico bruto (ex.: "Connection refused") subir até a interface.
        raise RuntimeError("Não foi possível falar com a ponte do WhatsApp. Verifique se o Vyzium e o WhatsApp Web estão abertos e tente novamente.") from exc

APP_NAME = "Vyzium"
DEFAULT_CONTROL_PRESETS = [
    {"id": "sent", "label": "Pedido enviado", "color": "#007D9C", "rule": "sent", "active": True},
    {"id": "card_payment", "label": "Pedido aguardando pagamento no cartão", "color": "#8756A5", "rule": "card_payment", "active": True},
    {"id": "cancel_requested", "label": "Fornecedor solicitou cancelamento", "color": "#B84E60", "rule": "cancel_requested", "active": True},
    {"id": "waiting_supplier", "label": "Aguardando retorno do fornecedor", "color": "#A77923", "rule": "waiting_supplier", "active": True},
    {"id": "delivery_scheduled", "label": "Entrega programada", "color": "#3578B9", "rule": "none", "active": True},
    {"id": "pending", "label": "Pendência", "color": "#B84E60", "rule": "none", "active": True},
    {"id": "request_collection", "label": "Solicitar coleta", "color": "#C79A20", "rule": "none", "active": True},
    {"id": "awaiting_collection", "label": "Aguardando ser coletado", "color": "#4C8792", "rule": "none", "active": True},
]


def approval_summary(items):
    states = set()
    for item in items:
        label = normalize(item.get("order_bpm_status", ""))
        if item.get("order_bpm_status_code") == 4 or any(word in label for word in ("RECUS", "REPROV", "REJEIT")):
            states.add("rejected")
        elif any(word in label for word in ("AGUARD", "PEND", "EM APROVAC", "NAO APROV")):
            states.add("waiting")
        elif item.get("order_bpm_status_code") == 3 or "APROVAD" in label:
            states.add("approved")
        else:
            states.add("unknown")
    status = next((x for x in ("rejected", "waiting", "unknown", "approved") if x in states), "unknown")
    return status, {"approved": "Aprovada", "waiting": "Aguardando aprovação", "rejected": "Recusada", "unknown": "Não informada"}[status]


def next_action(row, rule):
    if row["attendance_status"] in {"attended", "canceled"}:
        return "Sem ação pendente", "neutral"
    if rule == "cancel_requested":
        return "Validar cancelamento solicitado", "urgent"
    if row["approval_status"] == "rejected":
        return "Revisar motivo da recusa", "urgent"
    if row["approval_status"] == "waiting":
        return "Acompanhar aprovação da OC", "attention"
    if rule == "card_payment":
        return "Acompanhar pagamento no cartão", "finance"
    if rule == "sent" and row["urgency"] in {"due_soon", "overdue", "critical"}:
        return "Cobrar confirmação de entrega", "urgent" if row["urgency"] != "due_soon" else "attention"
    if row["urgency"] in {"overdue", "critical"} and not row["control_status"] and not row["control_note"].strip():
        return "Fornecedor ainda não atualizado", "urgent"
    if rule == "waiting_supplier":
        return "Cobrar retorno do fornecedor", "attention"
    if row["attendance_status"] == "partial":
        return "Cobrar saldo restante da OC", "attention"
    if not row["due_date"]:
        return "Solicitar previsão de entrega", "attention"
    if row["urgency"] in {"overdue", "critical"}:
        return "Atualizar previsão com fornecedor", "urgent"
    return "Acompanhar entrega", "normal"
CONTROL_STATUS_LABELS = {
    "": "Sem marcação",
    "sent": "Pedido enviado",
    "card_payment": "Aguardando pagamento no cartão",
    "waiting_supplier": "Aguardando retorno do fornecedor",
    "delivery_scheduled": "Entrega programada",
    "pending": "Pendência",
    "request_collection": "Solicitar coleta",
    "awaiting_collection": "Aguardando ser coletado",
}
ATTENDANCE_LABELS = {
    "pending": "Pendente",
    "partial": "Atendida parcialmente",
    "attended": "Atendida",
    "canceled": "Cancelada",
}
TARGET_FIELDS = {
    "oc": ["OC", "IDORDEMDECOMPRA", "ORDEM DE COMPRA", "NUMERO OC", "NUMERO DA OC"],
    "company_id": ["FKEMPRESA"],
    "company": ["EMPRESA", "HOTEL", "UNIDADE"],
    "purchase_type": ["TIPO DE COMPRA"],
    "sci": ["SCI", "IDSCI", "SOLICITACAO"],
    "sci_item_id": ["IDITEMDASCI"],
    "process_id": ["FKPROCESSODECOMPRA"],
    "article_code": ["CODIGOARTIGO", "CODIGO ARTIGO"],
    "quantity": ["QUANTIDADEOC", "QUANTIDADESCI", "QUANTIDADE"],
    "unit": ["UNIDADEMEDIDAOC", "FATOR OC", "FATOR", "UNIDADEMEDIDA", "UNIDADE DE MEDIDA"],
    "description": ["DESCRICAOARTIGO", "DESCRICAO DO ARTIGO", "ITEM", "DESCRICAO"],
    "buyer": ["COMPRADOR"],
    "buyer_id": ["FKCOMPRADOR"],
    "order_date": ["DATAOC", "DATA DA OC"],
    "due_date": ["DATAPREVISTAENTREGAOC", "DATA PREVISTA ENTREGA OC", "PREVISAO DE ENTREGA"],
    "received_date": ["DATAENTRADAMERCADORIA", "DATA ENTRADA MERCADORIA", "DATA DA ENTRADA"],
    "order_status": ["STATUSITEMDAORDEMDECOMPRA", "STATUS ITEM DA ORDEM DE COMPRA", "STATUS OC"],
    "order_status_code": ["NMSTATUSITEMDAORDEMDECOMPRA"],
    "order_bpm_status": ["STATUSBPMOC"],
    "order_bpm_status_code": ["NMSTATUSBPMOC"],
    "request_status": ["STATUSDOITEMDASCI", "STATUS DO ITEM DA SCI"],
    "supplier_id": ["FKFORNECEDOR"],
    "supplier_name": ["RAZAOSOCIALFORNECEDOR", "RAZAO SOCIAL FORNECEDOR", "NOMEFORNECEDOR", "FORNECEDOR"],
    "supplier_doc": ["CPFCNPJFORNECEDOR", "CPF CNPJ FORNECEDOR", "CNPJ"],
    "value_total": ["VALORTOTALITEMOC", "VALOR TOTAL ITEM OC", "VALOR TOTAL"],
    "value_unit": ["VALORUNITARIOITEMOC", "VALOR UNITARIO ITEM OC"],
    "received_qty": ["QUANTIDADERECEBIDA", "QUANTIDADE RECEBIDA"],
    "invoice": ["ANNUMERODANOTAFISCAL", "NUMERO DA NOTA FISCAL"],
    "receipt_unit": ["UNIDADEMEDIDARECEBIDA", "UNIDADE MEDIDA RECEBIDA"],
    "urgent": ["URGENTE"],
}
REQUIRED_FIELDS = {"oc", "company", "description", "due_date", "supplier_name"}


def normalize(value: Any) -> str:
    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    text = re.sub(r"[^A-Za-z0-9]+", " ", text.upper())
    return re.sub(r"\s+", " ", text).strip()


def person_matches(value: Any, filter_value: Any) -> bool:
    actual = set(normalize(value).split())
    expected = set(normalize(filter_value).split())
    return bool(actual and expected and (actual.issubset(expected) or expected.issubset(actual)))


def text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def iso_date(value: Any) -> str | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (int, float)):
        try:
            return from_excel(value).date().isoformat()
        except Exception:
            return None
    raw = str(value).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d-%m-%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return None


def number_value(value: Any) -> float | None:
    if value in (None, ""):
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def boolean_value(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if isinstance(value, (int, float)):
        return value != 0
    return normalize(value) in {"SIM", "S", "TRUE", "VERDADEIRO", "1", "YES", "Y"}


def phone_value(value: Any) -> str:
    """Normalize a phone to E.164-like form and repair legacy Brazilian mobiles.

    Brazilian mobile numbers use 2-digit DDD + 9-digit subscriber numbers. Some
    spreadsheets/old contact lists still contain the legacy 8-digit mobile form;
    when the subscriber starts with 6-9, the missing ninth digit is unambiguous
    enough to repair by inserting 9 after the DDD. Landlines (normally 2-5) are
    left untouched.
    """
    digits = re.sub(r"\D", "", str(value or ""))
    if digits.startswith("00"):
        digits = digits[2:]
    if digits and len(digits) in (10, 11):
        digits = "55" + digits
    if digits.startswith("55") and len(digits) == 12:
        subscriber = digits[4:]
        if subscriber and subscriber[0] in "6789":
            digits = digits[:4] + "9" + subscriber
    if digits and not 8 <= len(digits) <= 15:
        raise ValueError("Número de WhatsApp inválido. Confira país, DDD e quantidade de dígitos.")
    return "+" + digits if digits else ""


@dataclass(frozen=True)
class Classification:
    urgency: str
    days: int | None
    eligible: bool


def classify(due_date: str | None, received_date: str | None, order_status: str,
             request_status: str, warning_days: int = 3,
             critical_after_days: int = 10, today: date | None = None,
             order_status_code: int | None = None, received_qty: float | None = None,
             remaining_qty: float | None = None, quantity: float | None = None) -> Classification:
    status = normalize(f"{order_status} {request_status}")
    if order_status_code in {2, 3} or any(word in status for word in ("CANCELADO", "RECEBIDO TOTALMENTE", "ENTREGUE")):
        return Classification("completed", None, False)

    # A receipt date by itself does not prove that the whole item was received.
    # When quantity/saldo data exists, it is authoritative enough to distinguish
    # partial receipts from completion even if the numeric status is missing.
    if order_status_code is None and received_date and "PARCIAL" not in status:
        if remaining_qty is not None:
            if remaining_qty <= 1e-9 and (received_qty or 0) > 0:
                return Classification("completed", None, False)
        elif quantity is None or received_qty is None:
            return Classification("completed", None, False)
        elif received_qty >= quantity - 1e-9:
            return Classification("completed", None, False)
    if not due_date:
        return Classification("no_due_date", None, False)
    current = today or date.today()
    due = date.fromisoformat(due_date)
    delta = (due - current).days
    if delta > warning_days:
        return Classification("scheduled", delta, False)
    if delta >= 0:
        return Classification("due_soon", delta, True)
    late = abs(delta)
    if late > critical_after_days:
        return Classification("critical", -late, True)
    return Classification("overdue", -late, True)


class Store:
    def __init__(self, path: Path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.connection = sqlite3.connect(path, check_same_thread=False)
        self.connection.row_factory = sqlite3.Row
        self.lock = threading.RLock()
        with self.lock:
            self.connection.execute("PRAGMA journal_mode=WAL")
            self.connection.execute("PRAGMA foreign_keys=ON")
            self.connection.execute("PRAGMA busy_timeout=5000")
            self.connection.executescript("""
                CREATE TABLE IF NOT EXISTS orders (
                    item_key TEXT PRIMARY KEY,
                    oc TEXT NOT NULL,
                    company TEXT,
                    purchase_type TEXT,
                    sci TEXT,
                    quantity REAL,
                    unit TEXT,
                    description TEXT,
                    buyer TEXT,
                    due_date TEXT,
                    received_date TEXT,
                    order_status TEXT,
                    request_status TEXT,
                    supplier_key TEXT,
                    supplier_name TEXT,
                    supplier_doc TEXT,
                    value_total REAL,
                    company_id TEXT,
                    source_item_id TEXT,
                    process_id TEXT,
                    article_code TEXT,
                    buyer_id TEXT,
                    order_date TEXT,
                    order_status_code INTEGER,
                    order_bpm_status TEXT,
                    order_bpm_status_code INTEGER,
                    supplier_id TEXT,
                    value_unit REAL,
                    received_qty REAL DEFAULT 0,
                    remaining_qty REAL,
                    urgent INTEGER DEFAULT 0,
                    source_row_count INTEGER DEFAULT 1,
                    source_file TEXT,
                    imported_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_orders_due ON orders(due_date);
                CREATE INDEX IF NOT EXISTS idx_orders_supplier ON orders(supplier_key);
                CREATE INDEX IF NOT EXISTS idx_orders_oc_supplier ON orders(oc,supplier_key);
                CREATE INDEX IF NOT EXISTS idx_orders_buyer ON orders(buyer);
                CREATE INDEX IF NOT EXISTS idx_orders_company ON orders(company);
                CREATE TABLE IF NOT EXISTS suppliers (
                    supplier_key TEXT PRIMARY KEY,
                    display_name TEXT NOT NULL,
                    phone TEXT DEFAULT '',
                    contact_name TEXT DEFAULT '',
                    active INTEGER NOT NULL DEFAULT 1,
                    updated_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS receipts (
                    receipt_key TEXT PRIMARY KEY,
                    item_key TEXT NOT NULL,
                    receipt_date TEXT,
                    invoice TEXT,
                    quantity REAL,
                    unit TEXT
                );
                CREATE INDEX IF NOT EXISTS idx_receipts_item ON receipts(item_key);
                CREATE TABLE IF NOT EXISTS import_batches (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    source_file TEXT NOT NULL,
                    source_rows INTEGER NOT NULL,
                    order_items INTEGER NOT NULL,
                    duplicate_rows INTEGER NOT NULL,
                    imported_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS order_controls (
                    oc TEXT NOT NULL,
                    supplier_key TEXT NOT NULL,
                    control_status TEXT NOT NULL DEFAULT '',
                    note TEXT NOT NULL DEFAULT '',
                    sent_at TEXT,
                    updated_at TEXT NOT NULL,
                    PRIMARY KEY(oc, supplier_key)
                );
                CREATE TABLE IF NOT EXISTS order_control_history (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    oc TEXT NOT NULL,
                    supplier_key TEXT NOT NULL,
                    control_status TEXT NOT NULL DEFAULT '',
                    note TEXT NOT NULL DEFAULT '',
                    changed_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_order_control_history_order
                    ON order_control_history(oc, supplier_key, id DESC);
                CREATE TABLE IF NOT EXISTS followups (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL,
                    supplier_key TEXT NOT NULL,
                    supplier_name TEXT NOT NULL,
                    phone TEXT,
                    urgency TEXT NOT NULL,
                    message TEXT NOT NULL,
                    status TEXT NOT NULL,
                    error TEXT,
                    sent_at TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS followup_items (
                    followup_id INTEGER NOT NULL REFERENCES followups(id) ON DELETE CASCADE,
                    item_key TEXT NOT NULL,
                    urgency TEXT NOT NULL,
                    PRIMARY KEY(followup_id, item_key)
                );
                CREATE INDEX IF NOT EXISTS idx_followup_items_item ON followup_items(item_key,followup_id);
                CREATE TABLE IF NOT EXISTS message_batches (
                    batch_id TEXT PRIMARY KEY,
                    status TEXT NOT NULL,
                    simulation INTEGER NOT NULL DEFAULT 0,
                    request_json TEXT NOT NULL DEFAULT '{}',
                    selected_count INTEGER NOT NULL DEFAULT 0,
                    processed_count INTEGER NOT NULL DEFAULT 0,
                    sent_count INTEGER NOT NULL DEFAULT 0,
                    failed_count INTEGER NOT NULL DEFAULT 0,
                    uncertain_count INTEGER NOT NULL DEFAULT 0,
                    simulated_count INTEGER NOT NULL DEFAULT 0,
                    current_supplier TEXT,
                    error TEXT,
                    created_at TEXT NOT NULL,
                    started_at TEXT,
                    finished_at TEXT
                );
                CREATE INDEX IF NOT EXISTS idx_message_batches_status
                    ON message_batches(status, created_at);
                CREATE TABLE IF NOT EXISTS message_queue (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    batch_id TEXT NOT NULL REFERENCES message_batches(batch_id) ON DELETE CASCADE,
                    position INTEGER NOT NULL,
                    supplier_key TEXT NOT NULL,
                    supplier_name TEXT NOT NULL,
                    phone TEXT NOT NULL,
                    urgency TEXT NOT NULL,
                    message TEXT NOT NULL,
                    status TEXT NOT NULL DEFAULT 'queued',
                    attempts INTEGER NOT NULL DEFAULT 0,
                    followup_id INTEGER REFERENCES followups(id),
                    provider_message_id TEXT,
                    ack INTEGER,
                    error TEXT,
                    created_at TEXT NOT NULL,
                    started_at TEXT,
                    finished_at TEXT,
                    UNIQUE(batch_id, supplier_key)
                );
                CREATE INDEX IF NOT EXISTS idx_message_queue_batch
                    ON message_queue(batch_id, position);
                CREATE INDEX IF NOT EXISTS idx_message_queue_status
                    ON message_queue(status, batch_id, position);
                CREATE TABLE IF NOT EXISTS message_queue_items (
                    queue_id INTEGER NOT NULL REFERENCES message_queue(id) ON DELETE CASCADE,
                    item_key TEXT NOT NULL,
                    urgency TEXT NOT NULL,
                    PRIMARY KEY(queue_id, item_key)
                );
                CREATE TABLE IF NOT EXISTS settings (
                    key TEXT PRIMARY KEY,
                    value TEXT NOT NULL
                );
                CREATE TABLE IF NOT EXISTS app_logs (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    level TEXT NOT NULL,
                    event TEXT NOT NULL,
                    details TEXT,
                    created_at TEXT NOT NULL
                );
                CREATE INDEX IF NOT EXISTS idx_app_logs_created ON app_logs(id DESC);
            """)
            existing_order_columns = {row[1] for row in self.connection.execute("PRAGMA table_info(orders)")}
            order_migrations = {
                "company_id": "TEXT", "source_item_id": "TEXT", "process_id": "TEXT",
                "article_code": "TEXT", "buyer_id": "TEXT", "order_date": "TEXT",
                "order_status_code": "INTEGER", "order_bpm_status": "TEXT",
                "order_bpm_status_code": "INTEGER", "supplier_id": "TEXT", "value_unit": "REAL",
                "received_qty": "REAL DEFAULT 0", "remaining_qty": "REAL",
                "urgent": "INTEGER DEFAULT 0", "source_row_count": "INTEGER DEFAULT 1",
            }
            for column, definition in order_migrations.items():
                if column not in existing_order_columns:
                    self.connection.execute(f"ALTER TABLE orders ADD COLUMN {column} {definition}")
            defaults = {
                "warning_days": "3",
                "critical_after_days": "10",
                "cooldown_hours": "72",
                "simulation": "true",
                "automatic_enabled": "false",
                "schedule_time": "09:00",
                "last_workbook_path": "",
                "last_auto_run": "",
                "last_auto_attempt": "",
                "buyer_filter": "",
                "sender_name": "Compras",
                "message_signature": "Agradecemos desde já e aguardamos seu retorno."
            }
            self.connection.executemany(
                "INSERT OR IGNORE INTO settings(key,value) VALUES(?,?)", defaults.items()
            )
            self.connection.commit()
            self._repair_supplier_aliases()

    def _repair_supplier_aliases(self):
        """Merge supplier rows accidentally created with a normalized key.

        Versions prior to 1.2.5 normalized supplier_key while saving a phone.
        For document-based keys this could create a second supplier row detached
        from the orders. Only unambiguous aliases are merged.
        """
        order_keys = [row[0] for row in self.connection.execute(
            "SELECT DISTINCT supplier_key FROM orders WHERE supplier_key IS NOT NULL AND supplier_key<>''"
        )]
        by_normalized: dict[str, list[str]] = {}
        for key in order_keys:
            by_normalized.setdefault(normalize(key), []).append(key)
        aliases = [dict(row) for row in self.connection.execute("SELECT * FROM suppliers")]
        changed = False
        for alias in aliases:
            key = alias["supplier_key"]
            if key in order_keys:
                continue
            matches = by_normalized.get(normalize(key), [])
            if len(matches) != 1:
                continue
            canonical = matches[0]
            target = self.connection.execute(
                "SELECT * FROM suppliers WHERE supplier_key=?", (canonical,)
            ).fetchone()
            if not target:
                continue
            phone = target["phone"] or alias.get("phone", "")
            contact = target["contact_name"] or alias.get("contact_name", "")
            active = target["active"] if target["active"] is not None else alias.get("active", 1)
            self.connection.execute(
                "UPDATE suppliers SET phone=?,contact_name=?,active=?,updated_at=? WHERE supplier_key=?",
                (phone, contact, active, datetime.now().isoformat(timespec="seconds"), canonical),
            )
            self.connection.execute("DELETE FROM suppliers WHERE supplier_key=?", (key,))
            changed = True
        if changed:
            self.connection.commit()

    def log_event(self, event: str, details: Any = None, level: str = "info"):
        payload = None if details is None else json.dumps(details, ensure_ascii=False, default=str)
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            self.connection.execute(
                "INSERT INTO app_logs(level,event,details,created_at) VALUES(?,?,?,?)",
                (str(level)[:16], str(event)[:96], payload[:10000] if payload else None, now),
            )
            # Technical logs are diagnostic, not business history. Keep them bounded.
            self.connection.execute(
                "DELETE FROM app_logs WHERE id NOT IN (SELECT id FROM app_logs ORDER BY id DESC LIMIT 5000)"
            )
            self.connection.commit()

    def _prune_followups(self, keep: int = 20000):
        # Keep a large audit window while preventing an unattended installation
        # from growing forever. followup_items are removed by ON DELETE CASCADE.
        offset = max(0, int(keep) - 1)
        self.connection.execute(
            "DELETE FROM followups WHERE id < COALESCE((SELECT id FROM followups ORDER BY id DESC LIMIT 1 OFFSET ?), -1)",
            (offset,),
        )

    def current_order_count(self) -> int:
        with self.lock:
            return int(self.connection.execute("SELECT COUNT(*) FROM orders").fetchone()[0])

    @property
    def revision(self) -> int:
        """Monotonic in-process database revision used only for safe read caches.

        ``sqlite3.Connection.total_changes`` increments on every write made by
        this process, so cached operational views are automatically discarded
        whenever imports, controls, suppliers, settings or follow-up history change.
        """
        with self.lock:
            return int(self.connection.total_changes)

    def latest_import(self) -> dict[str, Any] | None:
        with self.lock:
            row = self.connection.execute(
                "SELECT * FROM import_batches ORDER BY id DESC LIMIT 1"
            ).fetchone()
        return dict(row) if row else None

    def settings(self) -> dict[str, Any]:
        with self.lock:
            raw = {row["key"]: row["value"] for row in self.connection.execute("SELECT key,value FROM settings")}
        presets = json.loads(raw.get("control_presets", json.dumps(DEFAULT_CONTROL_PRESETS)))
        # Novas marcações nativas também precisam aparecer para quem já possui
        # configurações salvas de versões anteriores.
        preset_ids = {preset.get("id") for preset in presets if isinstance(preset, dict)}
        presets.extend(dict(preset) for preset in DEFAULT_CONTROL_PRESETS if preset["id"] not in preset_ids)
        return {
            **raw,
            "control_presets": presets,
            "warning_days": int(raw.get("warning_days", 3)),
            "critical_after_days": int(raw.get("critical_after_days", 10)),
            "cooldown_hours": int(raw.get("cooldown_hours", 72)),
            "simulation": raw.get("simulation", "true").lower() == "true",
            "automatic_enabled": raw.get("automatic_enabled", "false").lower() == "true",
        }

    def save_settings(self, values: dict[str, Any]) -> dict[str, Any]:
        allowed = {"warning_days", "critical_after_days", "cooldown_hours", "simulation", "automatic_enabled", "schedule_time", "buyer_filter", "sender_name", "message_signature", "control_presets"}
        if "control_presets" in values:
            presets = values["control_presets"]
            if not isinstance(presets, list) or len(presets) > 200:
                raise ValueError("Lista de recomendações inválida (máximo 200).")
            ids = set()
            for preset in presets:
                if not isinstance(preset, dict) or not re.fullmatch(r"[a-zA-Z0-9_-]{1,64}", str(preset.get("id", ""))):
                    raise ValueError("Identificador de recomendação inválido.")
                if preset["id"] in ids or not str(preset.get("label", "")).strip() or len(str(preset["label"])) > 100:
                    raise ValueError("Nome ou identificador de recomendação inválido.")
                if not re.fullmatch(r"#[0-9a-fA-F]{6}", str(preset.get("color", ""))):
                    raise ValueError("Cor de recomendação inválida.")
                if preset.get("rule") not in {"none", "sent", "card_payment", "cancel_requested", "waiting_supplier"} or not isinstance(preset.get("active"), bool):
                    raise ValueError("Regra de recomendação inválida.")
                ids.add(preset["id"])
            values = {**values, "control_presets": json.dumps(presets, ensure_ascii=False)}
        numeric_limits = {
            "warning_days": (0, 30),
            "critical_after_days": (1, 365),
            "cooldown_hours": (1, 8760),
        }
        for key, (minimum, maximum) in numeric_limits.items():
            if key in values:
                try:
                    numeric = int(values[key])
                except (TypeError, ValueError) as exc:
                    raise ValueError(f"Configuração inválida: {key}.") from exc
                if not minimum <= numeric <= maximum:
                    raise ValueError(f"{key} deve estar entre {minimum} e {maximum}.")
                values[key] = numeric
        if "schedule_time" in values and not re.fullmatch(r"(?:[01]\d|2[0-3]):[0-5]\d", str(values["schedule_time"])):
            raise ValueError("O horário automático deve usar o formato HH:MM.")
        current = self.settings()
        prospective_simulation = bool(values.get("simulation", current["simulation"]))
        prospective_automatic = bool(values.get("automatic_enabled", current["automatic_enabled"]))
        if prospective_simulation and prospective_automatic:
            raise ValueError("Desative o modo simulação antes de ativar o envio automático.")
        with self.lock:
            for key, value in values.items():
                if key not in allowed:
                    continue
                stored = str(value).lower() if isinstance(value, bool) else str(value).strip()
                self.connection.execute(
                    "INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                    (key, stored),
                )
            self.connection.commit()
        return self.settings()

    def set_internal(self, key: str, value: str):
        if key not in {"last_workbook_path", "last_auto_run", "last_auto_attempt"}:
            raise ValueError("Configuração interna inválida.")
        with self.lock:
            self.connection.execute(
                "INSERT INTO settings(key,value) VALUES(?,?) ON CONFLICT(key) DO UPDATE SET value=excluded.value",
                (key, value),
            )
            self.connection.commit()

    def replace_snapshot(self, rows: list[dict[str, Any]], receipts: list[dict[str, Any]], source_file: str,
                         source_rows: int, duplicate_rows: int) -> dict[str, int]:
        if not rows:
            raise ValueError("A importação não encontrou itens de OC. A base anterior foi preservada.")
        previous = self.current_order_count()
        if previous >= 100 and len(rows) < previous * 0.5:
            raise ValueError("A nova base tem menos de 50% dos itens anteriores. A base anterior foi preservada para evitar perda acidental.")
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            try:
                self.connection.execute("BEGIN IMMEDIATE")
                existing_suppliers = {
                    normalize(row["display_name"]): dict(row)
                    for row in self.connection.execute("SELECT * FROM suppliers")
                }
                self.connection.execute("DELETE FROM receipts")
                self.connection.execute("DELETE FROM orders")
                columns = [
                    "item_key", "oc", "company", "purchase_type", "sci", "quantity", "unit", "description",
                    "buyer", "due_date", "received_date", "order_status", "request_status", "supplier_key",
                    "supplier_name", "supplier_doc", "value_total", "company_id", "source_item_id", "process_id",
                    "article_code", "buyer_id", "order_date", "order_status_code", "order_bpm_status",
                    "order_bpm_status_code", "supplier_id", "value_unit", "received_qty", "remaining_qty",
                    "urgent", "source_row_count"
                ]
                order_sql = f"INSERT INTO orders({','.join(columns)},source_file,imported_at) VALUES({','.join('?' for _ in columns)},?,?)"
                self.connection.executemany(
                    order_sql,
                    [tuple(row.get(column) for column in columns) + (source_file, now) for row in rows],
                )
                supplier_records: dict[str, tuple[Any, ...]] = {}
                for row in rows:
                    supplier_key = row["supplier_key"]
                    old = existing_suppliers.get(normalize(row["supplier_name"]))
                    phone = old.get("phone", "") if old else ""
                    contact = old.get("contact_name", "") if old else ""
                    supplier_records[supplier_key] = (supplier_key, row["supplier_name"], phone, contact, now)
                self.connection.executemany("""
                    INSERT INTO suppliers(supplier_key,display_name,phone,contact_name,active,updated_at)
                    VALUES(?,?,?,?,1,?) ON CONFLICT(supplier_key) DO UPDATE SET
                    display_name=excluded.display_name,updated_at=excluded.updated_at
                """, supplier_records.values())
                self.connection.executemany("""
                    INSERT INTO receipts(receipt_key,item_key,receipt_date,invoice,quantity,unit)
                    VALUES(:receipt_key,:item_key,:receipt_date,:invoice,:quantity,:unit)
                """, receipts)
                self.connection.execute("""
                    INSERT INTO import_batches(source_file,source_rows,order_items,duplicate_rows,imported_at)
                    VALUES(?,?,?,?,?)
                """, (source_file, source_rows, len(rows), duplicate_rows, now))
                self.connection.commit()
            except Exception:
                self.connection.rollback()
                raise
        return {"previous": previous, "order_items": len(rows), "receipts": len(receipts)}

    def suppliers(self) -> list[dict[str, Any]]:
        sql = """
            SELECT s.*, COUNT(o.item_key) AS order_items
            FROM suppliers s LEFT JOIN orders o ON o.supplier_key=s.supplier_key
            GROUP BY s.supplier_key ORDER BY s.display_name COLLATE NOCASE
        """
        with self.lock:
            return [dict(row) for row in self.connection.execute(sql)]

    def save_supplier(self, values: dict[str, Any]) -> dict[str, Any]:
        raw_key = values.get("supplier_key")
        supplier_key = text_value(raw_key) if raw_key not in (None, "") else normalize(values.get("display_name"))
        if not supplier_key:
            raise ValueError("Fornecedor inválido.")
        if len(supplier_key) > 256:
            raise ValueError("Identificador do fornecedor é muito grande.")
        with self.lock:
            existing = self.connection.execute(
                "SELECT * FROM suppliers WHERE supplier_key=?", (supplier_key,)
            ).fetchone()
            # Compatibility repair: if a caller still sends the old normalized
            # alias, prefer the unique key actually referenced by the orders.
            if not existing:
                matches = [row[0] for row in self.connection.execute(
                    "SELECT DISTINCT supplier_key FROM orders WHERE supplier_key IS NOT NULL"
                ) if normalize(row[0]) == normalize(supplier_key)]
                if len(matches) == 1:
                    supplier_key = matches[0]
                    existing = self.connection.execute(
                        "SELECT * FROM suppliers WHERE supplier_key=?", (supplier_key,)
                    ).fetchone()
        display_name = str(values.get("display_name") or (existing["display_name"] if existing else supplier_key)).strip()
        phone = phone_value(values.get("phone")) if "phone" in values else (existing["phone"] if existing else "")
        contact = str(values.get("contact_name") if "contact_name" in values else (existing["contact_name"] if existing else "")).strip()
        active = (1 if values.get("active") else 0) if "active" in values else (existing["active"] if existing else 1)
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            self.connection.execute("""
                INSERT INTO suppliers(supplier_key,display_name,phone,contact_name,active,updated_at)
                VALUES(?,?,?,?,?,?) ON CONFLICT(supplier_key) DO UPDATE SET
                display_name=excluded.display_name,phone=excluded.phone,contact_name=excluded.contact_name,
                active=excluded.active,updated_at=excluded.updated_at
            """, (supplier_key, display_name, phone, contact, active, now))
            self.connection.commit()
        self.log_event("supplier_saved", {"supplier_key": supplier_key, "phone_set": bool(phone), "active": bool(active)})
        return {"supplier_key": supplier_key, "display_name": display_name, "phone": phone, "contact_name": contact, "active": active}

    def order_rows(self) -> list[dict[str, Any]]:
        with self.lock:
            return [dict(row) for row in self.connection.execute("""
                SELECT o.*,s.phone,s.contact_name,s.active FROM orders o
                LEFT JOIN suppliers s ON s.supplier_key=o.supplier_key
                ORDER BY o.due_date,o.supplier_name,o.oc
            """)]

    def receipt_rows(self, item_keys: list[str]) -> list[dict[str, Any]]:
        if not item_keys:
            return []
        placeholders = ",".join("?" for _ in item_keys)
        with self.lock:
            return [dict(row) for row in self.connection.execute(
                f"SELECT * FROM receipts WHERE item_key IN ({placeholders}) ORDER BY receipt_date,invoice", item_keys
            )]

    def order_controls(self) -> dict[tuple[str, str], dict[str, Any]]:
        with self.lock:
            rows = self.connection.execute("SELECT * FROM order_controls").fetchall()
        return {(row["oc"], row["supplier_key"]): dict(row) for row in rows}

    def order_control(self, oc: str, supplier_key: str) -> dict[str, Any]:
        with self.lock:
            row = self.connection.execute(
                "SELECT * FROM order_controls WHERE oc=? AND supplier_key=?", (str(oc), supplier_key)
            ).fetchone()
            history = self.connection.execute("""
                SELECT control_status,note,changed_at FROM order_control_history
                WHERE oc=? AND supplier_key=? ORDER BY id DESC LIMIT 10
            """, (str(oc), supplier_key)).fetchall()
        current = dict(row) if row else {
            "oc": str(oc), "supplier_key": supplier_key, "control_status": "", "note": "",
            "sent_at": None, "updated_at": None,
        }
        labels = {**CONTROL_STATUS_LABELS, **{p["id"]: p["label"] for p in self.settings()["control_presets"]}}
        current["control_label"] = labels.get(current["control_status"], current["control_status"])
        current["history"] = [
            {**dict(item), "control_label": labels.get(item["control_status"], item["control_status"])}
            for item in history
        ]
        return current

    def save_order_control(self, values: dict[str, Any]) -> dict[str, Any]:
        oc = text_value(values.get("oc"))
        supplier_key = text_value(values.get("supplier_key"))
        control_status = text_value(values.get("control_status"))
        note = text_value(values.get("note"))
        if not oc or not supplier_key:
            raise ValueError("OC e fornecedor são obrigatórios.")
        if control_status and control_status not in {p["id"] for p in self.settings()["control_presets"]}:
            raise ValueError("Situação de controle inválida.")
        if len(note) > 1000:
            raise ValueError("A observação deve ter no máximo 1.000 caracteres.")
        with self.lock:
            exists = self.connection.execute(
                "SELECT 1 FROM orders WHERE oc=? AND supplier_key=? LIMIT 1", (oc, supplier_key)
            ).fetchone()
            if not exists:
                raise ValueError("A OC informada não existe na base atual.")
            previous = self.connection.execute(
                "SELECT * FROM order_controls WHERE oc=? AND supplier_key=?", (oc, supplier_key)
            ).fetchone()
            now = datetime.now().isoformat(timespec="seconds")
            sent_at = previous["sent_at"] if previous and control_status == "sent" else None
            if control_status == "sent" and (not previous or previous["control_status"] != "sent"):
                sent_at = now
            changed = not previous or previous["control_status"] != control_status or previous["note"] != note
            self.connection.execute("""
                INSERT INTO order_controls(oc,supplier_key,control_status,note,sent_at,updated_at)
                VALUES(?,?,?,?,?,?) ON CONFLICT(oc,supplier_key) DO UPDATE SET
                control_status=excluded.control_status,note=excluded.note,
                sent_at=excluded.sent_at,updated_at=excluded.updated_at
            """, (oc, supplier_key, control_status, note, sent_at, now))
            if changed:
                self.connection.execute("""
                    INSERT INTO order_control_history(oc,supplier_key,control_status,note,changed_at)
                    VALUES(?,?,?,?,?)
                """, (oc, supplier_key, control_status, note, now))
            self.connection.commit()
        return self.order_control(oc, supplier_key)

    def recent_history(self, limit: int = 100) -> list[dict[str, Any]]:
        with self.lock:
            return [dict(row) for row in self.connection.execute(
                "SELECT * FROM followups ORDER BY id DESC LIMIT ?", (limit,)
            )]

    def last_success(self, item_key: str) -> dict[str, Any] | None:
        with self.lock:
            row = self.connection.execute("""
                SELECT f.sent_at,fi.urgency,f.status FROM followup_items fi
                JOIN followups f ON f.id=fi.followup_id
                WHERE fi.item_key=? AND f.status IN ('sent','uncertain')
                ORDER BY f.id DESC LIMIT 1
            """, (item_key,)).fetchone()
        return dict(row) if row else None

    def last_successes(self, item_keys: list[str]) -> dict[str, dict[str, Any]]:
        """Return the latest sent/uncertain follow-up for many items in bulk.

        This replaces the historical N+1 lookup used while preparing WhatsApp
        groups without changing the anti-spam rules.
        """
        unique = list(dict.fromkeys(str(key) for key in item_keys if key))
        if not unique:
            return {}
        result: dict[str, dict[str, Any]] = {}
        # Stay below SQLite's conservative parameter limit on older builds.
        for offset in range(0, len(unique), 800):
            chunk = unique[offset:offset + 800]
            placeholders = ",".join("?" for _ in chunk)
            sql = f"""
                SELECT fi.item_key,f.sent_at,fi.urgency,f.status,f.id
                FROM followup_items fi
                JOIN followups f ON f.id=fi.followup_id
                WHERE fi.item_key IN ({placeholders})
                  AND f.status IN ('sent','uncertain')
                ORDER BY f.id DESC
            """
            with self.lock:
                rows = self.connection.execute(sql, chunk).fetchall()
            for row in rows:
                key = row["item_key"]
                if key not in result:
                    data = dict(row)
                    data.pop("id", None)
                    result[key] = data
        return result

    def finish_followup(self, followup_id, status, error=None):
        with self.lock:
            self.connection.execute("UPDATE followups SET status=?,error=? WHERE id=?", (status, error, followup_id))
            self.connection.commit()

    def review_followup(self, followup_id):
        with self.lock:
            self.connection.execute("UPDATE followups SET status='reviewed',error='Reenvio liberado após conferência manual.' WHERE id=? AND status='uncertain'", (int(followup_id),))
            self.connection.commit()
        return {"ok": True}

    def record_followup(self, batch_id: str, group: dict[str, Any], status: str, error: str | None = None) -> int:
        with self.lock:
            cursor = self.connection.execute("""
                INSERT INTO followups(batch_id,supplier_key,supplier_name,phone,urgency,message,status,error,sent_at)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (batch_id, group["supplier_key"], group["supplier_name"], group["phone"], group["urgency"],
                  group["message"], status, error, datetime.now().isoformat(timespec="seconds")))
            followup_id = cursor.lastrowid
            self.connection.executemany(
                "INSERT INTO followup_items(followup_id,item_key,urgency) VALUES(?,?,?)",
                [(followup_id, item["item_key"], item["urgency"]) for item in group["items"]]
            )
            self._prune_followups()
            self.connection.commit()
        return int(followup_id)


    def create_message_batch(self, batch_id: str, simulation: bool, request_payload: dict[str, Any]) -> dict[str, Any]:
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            self.connection.execute("""
                INSERT INTO message_batches(
                    batch_id,status,simulation,request_json,created_at
                ) VALUES(?,?,?,?,?)
            """, (batch_id, "preparing", int(bool(simulation)), json.dumps(request_payload, ensure_ascii=False), now))
            self.connection.commit()
        return self.message_batch(batch_id)

    def message_batch(self, batch_id: str) -> dict[str, Any] | None:
        with self.lock:
            row = self.connection.execute("SELECT * FROM message_batches WHERE batch_id=?", (batch_id,)).fetchone()
        if not row:
            return None
        data = dict(row)
        try:
            data["request"] = json.loads(data.pop("request_json") or "{}")
        except Exception:
            data["request"] = {}
        data["simulation"] = bool(data.get("simulation"))
        return data

    def recoverable_message_batch(self) -> dict[str, Any] | None:
        with self.lock:
            row = self.connection.execute("""
                SELECT * FROM message_batches
                WHERE status IN ('preparing','running','error')
                ORDER BY created_at DESC LIMIT 1
            """).fetchone()
        if not row:
            return None
        data = dict(row)
        try:
            data["request"] = json.loads(data.pop("request_json") or "{}")
        except Exception:
            data["request"] = {}
        data["simulation"] = bool(data.get("simulation"))
        return data

    def recover_interrupted_queue(self) -> int:
        """Conservatively quarantine work that was in-flight during a crash.

        A row in ``sending`` may already have reached WhatsApp. It must never be
        automatically replayed after a restart. Queued rows remain safe to resume.
        """
        now = datetime.now().isoformat(timespec="seconds")
        note = "Aplicativo interrompido durante o envio. Confira a conversa antes de liberar novo envio."
        with self.lock:
            rows = list(self.connection.execute(
                "SELECT id,followup_id,batch_id FROM message_queue WHERE status='sending'"
            ))
            for row in rows:
                self.connection.execute(
                    "UPDATE message_queue SET status='uncertain',error=?,finished_at=? WHERE id=?",
                    (note, now, row["id"])
                )
                if row["followup_id"]:
                    self.connection.execute(
                        "UPDATE followups SET status='uncertain',error=? WHERE id=?",
                        (note, row["followup_id"])
                    )
            self.connection.commit()
        return len(rows)

    def plan_message_batch(self, batch_id: str, groups: list[dict[str, Any]]) -> None:
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            existing = self.connection.execute(
                "SELECT COUNT(*) FROM message_queue WHERE batch_id=?", (batch_id,)
            ).fetchone()[0]
            if existing:
                return
            for position, group in enumerate(groups):
                cursor = self.connection.execute("""
                    INSERT INTO message_queue(
                        batch_id,position,supplier_key,supplier_name,phone,urgency,message,status,created_at
                    ) VALUES(?,?,?,?,?,?,?,?,?)
                """, (batch_id, position, group["supplier_key"], group["supplier_name"], group["phone"],
                      group["urgency"], group["message"], "queued", now))
                queue_id = int(cursor.lastrowid)
                self.connection.executemany(
                    "INSERT INTO message_queue_items(queue_id,item_key,urgency) VALUES(?,?,?)",
                    [(queue_id, item["item_key"], item["urgency"]) for item in group["items"]]
                )
            self.connection.execute("""
                UPDATE message_batches SET status='running',selected_count=?,started_at=COALESCE(started_at,?)
                WHERE batch_id=?
            """, (len(groups), now, batch_id))
            self.connection.commit()

    def queue_rows(self, batch_id: str) -> list[dict[str, Any]]:
        with self.lock:
            rows = [dict(row) for row in self.connection.execute(
                "SELECT * FROM message_queue WHERE batch_id=? ORDER BY position", (batch_id,)
            )]
            for row in rows:
                row["items"] = [dict(item) for item in self.connection.execute(
                    "SELECT item_key,urgency FROM message_queue_items WHERE queue_id=? ORDER BY item_key",
                    (row["id"],)
                )]
        return rows

    def claim_next_queue_item(self, batch_id: str) -> dict[str, Any] | None:
        """Atomically mark the next safe row as in-flight and create its guard record."""
        now = datetime.now().isoformat(timespec="seconds")
        guard_error = "Envio em andamento ou interrompido. Confira a conversa antes de reenviar."
        with self.lock:
            row = self.connection.execute("""
                SELECT * FROM message_queue
                WHERE batch_id=? AND status='queued'
                ORDER BY position LIMIT 1
            """, (batch_id,)).fetchone()
            if not row:
                return None
            data = dict(row)
            items = [dict(item) for item in self.connection.execute(
                "SELECT item_key,urgency FROM message_queue_items WHERE queue_id=? ORDER BY item_key",
                (data["id"],)
            )]
            cursor = self.connection.execute("""
                INSERT INTO followups(batch_id,supplier_key,supplier_name,phone,urgency,message,status,error,sent_at)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (batch_id, data["supplier_key"], data["supplier_name"], data["phone"], data["urgency"],
                  data["message"], "uncertain", guard_error, now))
            followup_id = int(cursor.lastrowid)
            self.connection.executemany(
                "INSERT INTO followup_items(followup_id,item_key,urgency) VALUES(?,?,?)",
                [(followup_id, item["item_key"], item["urgency"]) for item in items]
            )
            self.connection.execute("""
                UPDATE message_queue SET status='sending',attempts=attempts+1,followup_id=?,started_at=?,error=?
                WHERE id=? AND status='queued'
            """, (followup_id, now, guard_error, data["id"]))
            self.connection.execute(
                "UPDATE message_batches SET status='running',current_supplier=? WHERE batch_id=?",
                (data["supplier_name"], batch_id)
            )
            self._prune_followups()
            self.connection.commit()
            data.update({"status": "sending", "followup_id": followup_id, "items": items, "started_at": now})
            return data

    def finish_queue_item(self, queue_id: int, status: str, error: str | None = None,
                          provider_message_id: str | None = None, ack: int | None = None) -> None:
        if status not in {"sent", "failed", "uncertain", "simulated"}:
            raise ValueError("Status de fila inválido.")
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            row = self.connection.execute(
                "SELECT followup_id FROM message_queue WHERE id=?", (int(queue_id),)
            ).fetchone()
            self.connection.execute("""
                UPDATE message_queue SET status=?,error=?,provider_message_id=?,ack=?,finished_at=?
                WHERE id=?
            """, (status, error, provider_message_id, ack, now, int(queue_id)))
            if row and row["followup_id"]:
                self.connection.execute(
                    "UPDATE followups SET status=?,error=? WHERE id=?",
                    (status, error, row["followup_id"])
                )
            self.connection.commit()

    def finish_simulated_queue_item(self, batch_id: str, queue_id: int) -> int:
        """Simulation remains visible in history exactly like the previous flow."""
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            row = self.connection.execute("SELECT * FROM message_queue WHERE id=?", (int(queue_id),)).fetchone()
            items = list(self.connection.execute(
                "SELECT item_key,urgency FROM message_queue_items WHERE queue_id=?", (int(queue_id),)
            ))
            cursor = self.connection.execute("""
                INSERT INTO followups(batch_id,supplier_key,supplier_name,phone,urgency,message,status,error,sent_at)
                VALUES(?,?,?,?,?,?,?,?,?)
            """, (batch_id, row["supplier_key"], row["supplier_name"], row["phone"], row["urgency"],
                  row["message"], "simulated", None, now))
            followup_id = int(cursor.lastrowid)
            self.connection.executemany(
                "INSERT INTO followup_items(followup_id,item_key,urgency) VALUES(?,?,?)",
                [(followup_id, item["item_key"], item["urgency"]) for item in items]
            )
            self.connection.execute("""
                UPDATE message_queue SET status='simulated',followup_id=?,started_at=?,finished_at=?,error=NULL
                WHERE id=?
            """, (followup_id, now, now, int(queue_id)))
            self._prune_followups()
            self.connection.commit()
            return followup_id

    def refresh_message_batch(self, batch_id: str, *, finalize: bool = False, error: str | None = None) -> dict[str, Any]:
        now = datetime.now().isoformat(timespec="seconds")
        with self.lock:
            counts = {row["status"]: int(row["amount"]) for row in self.connection.execute("""
                SELECT status,COUNT(*) amount FROM message_queue WHERE batch_id=? GROUP BY status
            """, (batch_id,))}
            processed = sum(counts.get(key, 0) for key in ("sent", "failed", "uncertain", "simulated"))
            terminal = not any(counts.get(key, 0) for key in ("queued", "sending"))
            status = "done" if finalize or terminal else "running"
            self.connection.execute("""
                UPDATE message_batches SET status=?,processed_count=?,sent_count=?,failed_count=?,
                    uncertain_count=?,simulated_count=?,current_supplier=?,error=?,
                    finished_at=CASE WHEN ?='done' THEN ? ELSE finished_at END
                WHERE batch_id=?
            """, (status, processed, counts.get("sent",0), counts.get("failed",0), counts.get("uncertain",0),
                  counts.get("simulated",0), None if status == "done" else self.connection.execute(
                      "SELECT supplier_name FROM message_queue WHERE batch_id=? AND status='sending' LIMIT 1",
                      (batch_id,)
                  ).fetchone()[0] if counts.get("sending",0) else None, error, status, now, batch_id))
            self.connection.commit()
        return self.message_batch(batch_id) or {}

class WorkbookImporter:
    def __init__(self, store: Store):
        self.store = store
        self.lock = threading.Lock()

    @staticmethod
    def _field_map(values: list[Any]) -> dict[str, int]:
        headers = {normalize(value): index for index, value in enumerate(values) if value not in (None, "")}
        mapping = {}
        for field, aliases in TARGET_FIELDS.items():
            for alias in aliases:
                if normalize(alias) in headers:
                    mapping[field] = headers[normalize(alias)]
                    break
        return mapping

    def _find_source(self, workbook) -> tuple[Any, int, dict[str, int]]:
        best = None
        for sheet in workbook.worksheets:
            for row_index, values in enumerate(sheet.iter_rows(min_row=1, max_row=25, values_only=True), 1):
                mapping = self._field_map(list(values))
                score = len(mapping) + 10 * len(REQUIRED_FIELDS.intersection(mapping))
                if best is None or score > best[0]:
                    best = (score, sheet, row_index, mapping)
        if not best or not REQUIRED_FIELDS.issubset(best[3]):
            raise ValueError(
                "Não encontrei uma aba de pedidos válida. Selecione a BASE SCI.xlsx ou a planilha completa "
                "de acompanhamento. A aba precisa conter: OC (ou IDORDEMDECOMPRA), EMPRESA, "
                "DESCRICAOARTIGO, DATAPREVISTAENTREGAOC e RAZAOSOCIALFORNECEDOR. "
                "A base anterior foi preservada."
            )
        return best[1], best[2], best[3]

    def import_file(self, path_value: str) -> dict[str, Any]:
        with self.lock:
            return self._import_file(path_value)

    def _import_file(self, path_value: str) -> dict[str, Any]:
        path = Path(path_value).expanduser().resolve()
        if path.suffix.lower() not in {".xlsx", ".xlsm"} or not path.is_file():
            raise ValueError("Selecione um arquivo .xlsx ou .xlsm válido.")
        workbook = load_workbook(path, read_only=True, data_only=False, keep_vba=path.suffix.lower() == ".xlsm")
        try:
            for candidate in workbook.worksheets:
                if candidate.max_column in (None, 1):
                    candidate.reset_dimensions()
            sheet, header_row, mapping = self._find_source(workbook)
            groups: dict[str, dict[str, Any]] = {}
            skipped = 0
            source_rows = 0
            duplicate_rows = 0
            seen_rows: set[tuple[Any, ...]] = set()
            for values in sheet.iter_rows(min_row=header_row + 1, values_only=True):
                source_rows += 1
                signature = tuple(values)
                if signature in seen_rows:
                    duplicate_rows += 1
                    continue
                seen_rows.add(signature)

                def get(field):
                    index = mapping.get(field)
                    return values[index] if index is not None and index < len(values) else None

                oc = text_value(get("oc"))
                supplier_name = text_value(get("supplier_name"))
                description = text_value(get("description"))
                if not oc or not supplier_name or not description:
                    skipped += 1
                    continue
                source_item_id = text_value(get("sci_item_id"))
                supplier_id = text_value(get("supplier_id"))
                supplier_doc = text_value(get("supplier_doc"))
                supplier_key = supplier_id or supplier_doc or normalize(supplier_name)
                source_identity = source_item_id or "|".join([
                    text_value(get("sci")), text_value(get("article_code")), normalize(description)
                ])
                identity = "|".join([source_identity, oc, supplier_key])
                item_key = hashlib.sha256(identity.encode("utf-8")).hexdigest()[:24]
                row = groups.setdefault(item_key, {
                    "item_key": item_key,
                    "oc": oc,
                    "company_id": text_value(get("company_id")),
                    "company": text_value(get("company")),
                    "purchase_type": text_value(get("purchase_type")),
                    "sci": text_value(get("sci")),
                    "source_item_id": source_item_id,
                    "process_id": text_value(get("process_id")),
                    "article_code": text_value(get("article_code")),
                    "quantity": number_value(get("quantity")),
                    "unit": text_value(get("unit")),
                    "description": description,
                    "buyer_id": text_value(get("buyer_id")),
                    "buyer": text_value(get("buyer")),
                    "order_date": iso_date(get("order_date")),
                    "due_date": iso_date(get("due_date")),
                    "received_date": iso_date(get("received_date")),
                    "order_status": text_value(get("order_status")),
                    "order_status_code": int(number_value(get("order_status_code"))) if number_value(get("order_status_code")) is not None else None,
                    "order_bpm_status": text_value(get("order_bpm_status")),
                    "order_bpm_status_code": int(number_value(get("order_bpm_status_code"))) if number_value(get("order_bpm_status_code")) is not None else None,
                    "request_status": text_value(get("request_status")),
                    "supplier_key": supplier_key,
                    "supplier_id": supplier_id,
                    "supplier_name": supplier_name,
                    "supplier_doc": supplier_doc,
                    "value_unit": number_value(get("value_unit")),
                    "value_total": number_value(get("value_total")),
                    "received_qty": 0.0,
                    "remaining_qty": number_value(get("quantity")),
                    "urgent": 1 if boolean_value(get("urgent")) else 0,
                    "source_row_count": 0,
                    "_receipts": {},
                })
                row["source_row_count"] += 1
                if boolean_value(get("urgent")):
                    row["urgent"] = 1
                received_date = iso_date(get("received_date"))
                if received_date and (not row.get("received_date") or received_date > row["received_date"]):
                    row["received_date"] = received_date
                received_qty = number_value(get("received_qty"))
                if received_qty is not None:
                    invoice = text_value(get("invoice"))
                    receipt_unit = text_value(get("receipt_unit"))
                    receipt_identity = "|".join([item_key, received_date or "", invoice, str(received_qty), receipt_unit])
                    receipt_key = hashlib.sha256(receipt_identity.encode("utf-8")).hexdigest()[:24]
                    row["_receipts"][receipt_key] = {
                        "receipt_key": receipt_key, "item_key": item_key, "receipt_date": received_date,
                        "invoice": invoice, "quantity": received_qty, "unit": receipt_unit,
                    }
            rows = []
            receipts = []
            for row in groups.values():
                item_receipts = list(row.pop("_receipts").values())
                row["received_qty"] = sum(receipt["quantity"] or 0 for receipt in item_receipts)
                ordered = row.get("quantity")
                row["remaining_qty"] = max(0.0, ordered - row["received_qty"]) if ordered is not None else None
                rows.append(row)
                receipts.extend(item_receipts)
            result = self.store.replace_snapshot(rows, receipts, path.name, source_rows, duplicate_rows)
            self.store.set_internal("last_workbook_path", str(path))
            return {
                **result,
                "read": source_rows,
                "order_items": len(rows),
                "receipts": len(receipts),
                "duplicate_rows": duplicate_rows,
                "skipped": skipped,
                "sheet": sheet.title,
                "header_row": header_row,
                "mapped_fields": sorted(mapping),
                "source": path.name,
            }
        finally:
            workbook.close()


URGENCY_RANK = {"due_soon": 1, "overdue": 2, "critical": 3}
URGENCY_LABEL = {
    "due_soon": "Próximo do prazo",
    "overdue": "Atrasado",
    "critical": "Atraso crítico",
    "scheduled": "Programado",
    "completed": "Concluído",
    "no_due_date": "Sem previsão",
}


class FollowUpService:
    def __init__(self, store: Store):
        self.store = store
        self.cache_lock = threading.RLock()
        self.cache_revision = -1
        self.orders_cache: list[dict[str, Any]] | None = None
        self.summary_cache: dict[tuple[str, str, str, str, str, str], list[dict[str, Any]]] = {}
        self.send_lock = threading.Lock()
        self.job_lock = threading.Lock()
        self.worker_batches: set[str] = set()
        self.store.recover_interrupted_queue()
        self.job_state: dict[str, Any] = {
            "active": False, "job_id": None, "phase": "idle", "total": 0,
            "processed": 0, "sent": 0, "failed": 0, "uncertain": 0,
            "simulated": 0, "current_supplier": None, "error": None, "result": None,
        }
        recoverable = self.store.recoverable_message_batch()
        if recoverable:
            self.job_state = self._state_from_batch(recoverable, active=True)
            # The WhatsApp bridge already exists before the Python engine starts.
            # Resume only queued work; any item that was in-flight was quarantined
            # as uncertain by recover_interrupted_queue() and is never replayed.
            self._start_persistent_worker(recoverable["batch_id"])

    def _sync_read_cache(self) -> int:
        """Invalidate derived read caches after any database write."""
        revision = self.store.revision
        with self.cache_lock:
            if revision != self.cache_revision:
                self.cache_revision = revision
                self.orders_cache = None
                self.summary_cache.clear()
        return revision

    def _update_job(self, **changes):
        with self.job_lock:
            self.job_state.update(changes)

    def _state_from_batch(self, batch: dict[str, Any], active: bool | None = None) -> dict[str, Any]:
        status = batch.get("status", "preparing")
        if active is None:
            active = status in {"preparing", "running"}
        result = None
        if not active and status == "done":
            result = self._batch_result(batch["batch_id"])
        return {
            "active": bool(active),
            "job_id": batch.get("batch_id"),
            "phase": "preparing" if status == "preparing" else ("done" if status == "done" else ("error" if status == "error" else "sending")),
            "total": int(batch.get("selected_count") or 0),
            "processed": int(batch.get("processed_count") or 0),
            "sent": int(batch.get("sent_count") or 0),
            "failed": int(batch.get("failed_count") or 0),
            "uncertain": int(batch.get("uncertain_count") or 0),
            "simulated": int(batch.get("simulated_count") or 0),
            "current_supplier": batch.get("current_supplier"),
            "error": batch.get("error"),
            "result": result,
            "started_at": batch.get("started_at") or batch.get("created_at"),
            "finished_at": batch.get("finished_at"),
        }

    def _batch_result(self, batch_id: str) -> dict[str, Any]:
        batch = self.store.message_batch(batch_id) or {}
        rows = self.store.queue_rows(batch_id)
        results = [{
            "id": row.get("followup_id"),
            "supplier_name": row.get("supplier_name"),
            "status": row.get("status"),
            "error": row.get("error"),
        } for row in rows]
        return {
            "batch_id": batch_id,
            "sent": int(batch.get("sent_count") or 0),
            "simulated": int(batch.get("simulated_count") or 0),
            "failed": int(batch.get("failed_count") or 0),
            "processed": int(batch.get("processed_count") or 0),
            "selected": int(batch.get("selected_count") or len(rows)),
            "uncertain": int(batch.get("uncertain_count") or 0),
            "aborted": False, "aborted_error": None,
            "results": results,
        }

    def send_status(self, job_id: str | None = None) -> dict[str, Any]:
        with self.job_lock:
            state = dict(self.job_state)
            if isinstance(state.get("result"), dict):
                state["result"] = dict(state["result"])
        persisted = self.store.message_batch(job_id) if job_id else None
        if job_id and persisted:
            with self.job_lock:
                active = job_id in self.worker_batches or (
                    self.job_state.get("active") and self.job_state.get("job_id") == job_id
                )
            state = self._state_from_batch(persisted, active=active)
            if not active and persisted.get("status") == "done":
                state["result"] = self._batch_result(job_id)
            return state
        if job_id and state.get("job_id") and job_id != state.get("job_id"):
            raise ValueError("Lote não encontrado.")
        return state

    @staticmethod
    def _validate_selection(supplier_keys, message_overrides):
        if supplier_keys is None:
            selected = None
        else:
            if not isinstance(supplier_keys, list) or len(supplier_keys) > 500:
                raise ValueError("Seleção de fornecedores inválida.")
            cleaned = [str(key).strip() for key in supplier_keys if str(key).strip()]
            if not cleaned:
                raise ValueError("Nenhum fornecedor selecionado. O envio foi cancelado por segurança.")
            selected = cleaned

        overrides: dict[str, str] = {}
        if message_overrides is not None:
            if not isinstance(message_overrides, dict) or len(message_overrides) > 200:
                raise ValueError("Edições de mensagem inválidas.")
            for raw_key, raw_message in message_overrides.items():
                key = str(raw_key).strip()
                if not key or len(key) > 256 or not isinstance(raw_message, str):
                    raise ValueError("Edição de mensagem inválida.")
                message = raw_message.strip()
                if not message:
                    raise ValueError("A mensagem editada não pode ficar vazia.")
                if len(message) > 60000:
                    raise ValueError("A mensagem editada é muito grande.")
                overrides[key] = message
        return selected, overrides

    def _groups_for_buyer(self, buyer: str = "", include_blocked: bool = False) -> list[dict[str, Any]]:
        buyer = str(buyer or "").strip()
        # Keep compatibility with tests/custom integrations that monkey-patch
        # groups() using the historical zero-argument signature.
        if not buyer and not include_blocked:
            return self.groups()
        return self.groups(include_blocked=include_blocked, buyer=buyer)

    def _groups_for_request(self, request: dict[str, Any], simulation: bool) -> list[dict[str, Any]]:
        selected_list, overrides = self._validate_selection(
            request.get("supplier_keys"), request.get("message_overrides")
        )
        selected = set(selected_list) if selected_list is not None else None
        buyer = str(request.get("buyer") or "").strip()
        groups = [group for group in self._groups_for_buyer(buyer) if selected is None or group["supplier_key"] in selected]
        if groups and not simulation:
            self._wait_for_connection()
            # Same established rule as before: after connection becomes ready,
            # re-read eligibility so a stale snapshot can never be sent.
            groups = [group for group in self._groups_for_buyer(buyer) if selected is None or group["supplier_key"] in selected]
        for group in groups:
            edited = overrides.get(group["supplier_key"])
            if edited is not None:
                group["message"] = edited
        return groups

    def _wait_for_connection(self):
        """Wait silently while the connection manager restores WhatsApp.

        Pausing the connection does not destroy the queue. The worker simply waits
        until the user explicitly resumes it. This method emits no UI notification;
        the existing WhatsApp panel remains the only visible status surface.
        """
        while True:
            try:
                return whatsapp_request("/wait")
            except RuntimeError:
                time.sleep(3)

    def _start_persistent_worker(self, batch_id: str):
        with self.job_lock:
            if batch_id in self.worker_batches:
                return
            self.worker_batches.add(batch_id)
            self.job_state["active"] = True
            self.job_state["job_id"] = batch_id

        def worker():
            last_error = None
            try:
                for attempt in range(5):
                    try:
                        self._run_persistent_batch(batch_id)
                        last_error = None
                        break
                    except Exception as exc:
                        last_error = str(exc)
                        # Queue rows remain persisted. A transient engine/bridge
                        # problem retries in-place without losing the job. In-flight
                        # rows are guarded by the followup record and are never replayed.
                        with self.store.lock:
                            self.store.connection.execute(
                                "UPDATE message_batches SET status='running',error=? WHERE batch_id=?",
                                (last_error, batch_id)
                            )
                            self.store.connection.commit()
                        self._update_job(active=True, error=None, current_supplier=None)
                        time.sleep(min(30, 3 * (2 ** attempt)))
                if last_error:
                    with self.store.lock:
                        now = datetime.now().isoformat(timespec="seconds")
                        self.store.connection.execute(
                            "UPDATE message_batches SET status='error',error=?,current_supplier=NULL,finished_at=? WHERE batch_id=?",
                            (last_error, now, batch_id)
                        )
                        self.store.connection.commit()
            finally:
                with self.job_lock:
                    self.worker_batches.discard(batch_id)
                batch = self.store.message_batch(batch_id)
                if batch and batch.get("status") == "done":
                    self._update_job(**self._state_from_batch(batch, active=False))
                    self._update_job(result=self._batch_result(batch_id))
                elif batch:
                    self._update_job(**self._state_from_batch(batch, active=False))

        threading.Thread(target=worker, name=f"followup-queue-{batch_id[:8]}", daemon=True).start()

    def start_send(self, supplier_keys=None, force_simulation=None, message_overrides=None, buyer: str = "") -> dict[str, Any]:
        selected, overrides = self._validate_selection(supplier_keys, message_overrides)
        settings = self.store.settings()
        simulation = settings["simulation"] if force_simulation is None else bool(force_simulation)
        with self.job_lock:
            if self.job_state.get("active") or self.send_lock.locked() or self.worker_batches:
                raise ValueError("Já existe um lote em execução ou aguardando conexão.")
            job_id = str(uuid.uuid4())
            request_payload = {"supplier_keys": selected, "message_overrides": overrides, "buyer": str(buyer or "").strip()}
            batch = self.store.create_message_batch(job_id, simulation, request_payload)
            self.job_state = self._state_from_batch(batch, active=True)
        self._start_persistent_worker(job_id)
        return self.send_status(job_id)

    def _run_persistent_batch(self, batch_id: str) -> dict[str, Any]:
        if not self.send_lock.acquire(blocking=False):
            raise ValueError("Já existe um lote em execução ou aguardando conexão.")
        try:
            batch = self.store.message_batch(batch_id)
            if not batch:
                raise ValueError("Lote não encontrado.")
            simulation = bool(batch.get("simulation"))
            rows = self.store.queue_rows(batch_id)
            if not rows:
                groups = self._groups_for_request(batch.get("request") or {}, simulation)
                self.store.plan_message_batch(batch_id, groups)
                rows = self.store.queue_rows(batch_id)
            batch = self.store.refresh_message_batch(batch_id)
            self._update_job(**self._state_from_batch(batch, active=True))

            while True:
                rows = self.store.queue_rows(batch_id)
                queued = next((row for row in rows if row.get("status") == "queued"), None)
                if not queued:
                    break

                if simulation:
                    self._update_job(current_supplier=queued.get("supplier_name"), phase="sending")
                    self.store.finish_simulated_queue_item(batch_id, queued["id"])
                    batch = self.store.refresh_message_batch(batch_id)
                    self._update_job(**self._state_from_batch(batch, active=True))
                    continue

                # Verify the session immediately before every provider. The Node
                # connection manager performs the actual health checks/reconnects.
                self._wait_for_connection()
                item = self.store.claim_next_queue_item(batch_id)
                if not item:
                    continue
                self._update_job(current_supplier=item.get("supplier_name"), phase="sending")
                status, error, result = "uncertain", None, {}
                try:
                    result = whatsapp_request("/send", {"phone": item["phone"], "message": item["message"]})
                    status = result.get("status", "uncertain")
                    if status not in {"sent", "failed", "uncertain"}:
                        status = "uncertain"
                    error = result.get("error")
                except Exception:
                    status = "uncertain"
                    error = "Resposta interrompida. Confira a conversa no WhatsApp antes de liberar outro envio."

                self.store.finish_queue_item(
                    item["id"], status, error,
                    provider_message_id=result.get("message_id"), ack=result.get("ack")
                )
                if status == "sent":
                    resolved_phone = result.get("resolved_phone")
                    if resolved_phone and resolved_phone != item.get("phone"):
                        self.store.save_supplier({
                            "supplier_key": item["supplier_key"],
                            "display_name": item["supplier_name"],
                            "phone": resolved_phone,
                        })
                    time.sleep(2)
                self.store.log_event("followup_result", {
                    "batch_id": batch_id, "supplier_key": item["supplier_key"],
                    "status": status, "error": error, "simulation": False
                }, "error" if status in {"failed", "uncertain"} else "info")
                batch = self.store.refresh_message_batch(batch_id)
                self._update_job(**self._state_from_batch(batch, active=True))

            batch = self.store.refresh_message_batch(batch_id, finalize=True)
            result = self._batch_result(batch_id)
            self._update_job(**self._state_from_batch(batch, active=False))
            self._update_job(result=result)
            return result
        finally:
            self.send_lock.release()

    def send_persistent_and_wait(self, supplier_keys=None, force_simulation=None, message_overrides=None, buyer: str = "") -> dict[str, Any]:
        state = self.start_send(supplier_keys, force_simulation, message_overrides, buyer)
        job_id = state["job_id"]
        while state.get("active"):
            time.sleep(0.2)
            state = self.send_status(job_id)
        if state.get("phase") == "error":
            raise RuntimeError(state.get("error") or "O lote foi interrompido.")
        return state.get("result") or self._batch_result(job_id)

    def orders(self, status: str | None = None) -> list[dict[str, Any]]:
        self._sync_read_cache()
        with self.cache_lock:
            cached = self.orders_cache
        if cached is None:
            settings = self.store.settings()
            built: list[dict[str, Any]] = []
            for source in self.store.order_rows():
                row = dict(source)
                state = classify(row["due_date"], row["received_date"], row["order_status"], row["request_status"],
                                 settings["warning_days"], settings["critical_after_days"],
                                 order_status_code=row.get("order_status_code"),
                                 received_qty=row.get("received_qty"), remaining_qty=row.get("remaining_qty"),
                                 quantity=row.get("quantity"))
                data_warning = ""
                if row.get("order_status_code") == 1 and (row.get("remaining_qty") or 0) <= 0:
                    data_warning = "Status parcial, mas a quantidade recebida já cobre o pedido. Confira a base."
                row.update({"urgency": state.urgency, "urgency_label": URGENCY_LABEL[state.urgency], "days": state.days,
                            "eligible": state.eligible, "data_warning": data_warning})
                built.append(row)
            # Cache only if no write happened while the snapshot was being built.
            if self.store.revision == self.cache_revision:
                with self.cache_lock:
                    self.orders_cache = built
            cached = built
        if status and status != "all":
            return [dict(row) for row in cached if row["urgency"] == status]
        return [dict(row) for row in cached]

    def filters(self) -> dict[str, list[str]]:
        rows = self.orders()
        return {
            "buyers": sorted({row["buyer"] for row in rows if row.get("buyer")}),
            "companies": sorted({row["company"] for row in rows if row.get("company")}),
            "suppliers": sorted({row["supplier_name"] for row in rows if row.get("supplier_name")}),
        }

    @staticmethod
    def _summary_urgency(items: list[dict[str, Any]]) -> str:
        open_items = [item for item in items if item["urgency"] not in {"completed"}]
        candidates = open_items or items
        rank = {"critical": 5, "overdue": 4, "due_soon": 3, "scheduled": 2, "no_due_date": 1, "completed": 0}
        return max((item["urgency"] for item in candidates), key=lambda value: rank[value])

    @staticmethod
    def _item_attendance(item: dict[str, Any]) -> str:
        code = item.get("order_status_code")
        if code == 3:
            return "canceled"
        if code == 2:
            return "attended"
        if code == 1:
            return "partial"
        if code == 0:
            return "pending"
        status = normalize(item.get("order_status"))
        if "CANCEL" in status:
            return "canceled"
        if "PARCIAL" in status:
            return "partial"
        if "TOTAL" in status or "ENTREGUE" in status:
            return "attended"
        received = item.get("received_qty")
        remaining = item.get("remaining_qty")
        if remaining is not None and (received or 0) > 0:
            return "attended" if remaining <= 1e-9 else "partial"
        if item.get("received_date") and remaining is not None and remaining > 1e-9:
            return "partial"
        if item.get("received_date") and remaining is not None and remaining <= 1e-9:
            return "attended"
        return "pending"

    @classmethod
    def _summary_attendance(cls, items: list[dict[str, Any]]) -> str:
        states = [cls._item_attendance(item) for item in items]
        active_states = [state for state in states if state != "canceled"]
        if not active_states:
            return "canceled"
        if all(state == "attended" for state in active_states):
            return "attended"
        if "partial" in active_states or ("attended" in active_states and "pending" in active_states):
            return "partial"
        return "pending"

    def order_summaries(self, buyer: str = "", company: str = "", urgency: str = "", search: str = "",
                        control_status: str = "all", attendance_status: str = "all") -> list[dict[str, Any]]:
        self._sync_read_cache()
        cache_key = (str(buyer), str(company), str(urgency), str(search), str(control_status), str(attendance_status))
        with self.cache_lock:
            cached = self.summary_cache.get(cache_key)
        if cached is not None:
            return [dict(row) for row in cached]
        grouped: dict[tuple[str, str], list[dict[str, Any]]] = {}
        buyer_key, company_key, search_key = normalize(buyer), normalize(company), normalize(search)
        controls = self.store.order_controls()
        presets = {p["id"]: p for p in self.store.settings()["control_presets"]}
        for item in self.orders():
            if buyer_key and not person_matches(item.get("buyer"), buyer):
                continue
            if company_key and normalize(item.get("company")) != company_key:
                continue
            grouped.setdefault((item["oc"], item["supplier_key"]), []).append(item)
        summaries = []
        for (oc, supplier_key), items in grouped.items():
            state = self._summary_urgency(items)
            attendance = self._summary_attendance(items)
            if urgency == "open" and state == "completed":
                continue
            if urgency and urgency not in {"all", "open"} and state != urgency:
                continue
            if attendance_status not in {"", "all"} and attendance != attendance_status:
                continue
            first = items[0]
            control = controls.get((oc, supplier_key), {})
            current_control = control.get("control_status", "")
            if control_status == "blank" and current_control:
                continue
            if control_status not in {"", "all", "blank"} and current_control != control_status:
                continue
            preset = presets.get(current_control, {})
            control_label = preset.get("label", CONTROL_STATUS_LABELS.get(current_control, current_control))
            haystack = normalize(" ".join([
                oc, first.get("supplier_name", ""), first.get("company", ""), first.get("buyer", ""),
                control_label, control.get("note", ""), *[x.get("description", "") for x in items]
            ]))
            if search_key and search_key not in haystack:
                continue
            open_items = [item for item in items if item["urgency"] != "completed"]
            due_dates = [item["due_date"] for item in (open_items or items) if item.get("due_date")]
            total_value = sum(item.get("value_total") or 0 for item in items)
            summaries.append({
                "oc": oc, "supplier_key": supplier_key, "supplier_name": first.get("supplier_name", ""),
                "supplier_doc": first.get("supplier_doc", ""), "company": first.get("company", ""),
                "buyer": first.get("buyer", ""), "order_date": first.get("order_date"),
                "due_date": min(due_dates) if due_dates else None, "urgency": state,
                "urgency_label": URGENCY_LABEL[state], "item_count": len(items),
                "attendance_status": attendance, "attendance_label": ATTENDANCE_LABELS[attendance],
                "open_item_count": len(open_items), "partial_item_count": sum(1 for item in items if item.get("order_status_code") == 1),
                "total_value": total_value, "urgent": any(bool(item.get("urgent")) for item in items),
                "warning_count": sum(1 for item in items if item.get("data_warning")),
                "control_status": current_control, "control_label": control_label,
                "control_note": control.get("note", ""), "control_sent_at": control.get("sent_at"),
                "control_updated_at": control.get("updated_at"),
                "phone": first.get("phone") or "",
            })
            row = summaries[-1]
            row["control_color"] = preset.get("color", "#647789")
            row["approval_status"], row["approval_label"] = approval_summary(items)
            row["next_action"], row["action_priority"] = next_action(row, preset.get("rule", "none"))
        rank = {"critical": 0, "overdue": 1, "due_soon": 2, "scheduled": 3, "no_due_date": 4, "completed": 5}
        result = sorted(summaries, key=lambda row: (rank[row["urgency"]], row.get("due_date") or "9999", row["oc"]))
        if self.store.revision == self.cache_revision:
            with self.cache_lock:
                if len(self.summary_cache) >= 32:
                    self.summary_cache.clear()
                self.summary_cache[cache_key] = result
        return [dict(row) for row in result]

    def order_detail(self, oc: str, supplier_key: str = "") -> dict[str, Any]:
        items = [item for item in self.orders() if item["oc"] == str(oc) and (not supplier_key or item["supplier_key"] == supplier_key)]
        if not items:
            raise ValueError("OC não encontrada.")
        receipts = self.store.receipt_rows([item["item_key"] for item in items])
        by_item: dict[str, list[dict[str, Any]]] = {}
        for receipt in receipts:
            by_item.setdefault(receipt["item_key"], []).append(receipt)
        for item in items:
            item["receipts"] = by_item.get(item["item_key"], [])
        summary = self.order_summaries(search=str(oc))
        summary = next((row for row in summary if row["oc"] == str(oc) and (not supplier_key or row["supplier_key"] == supplier_key)), None)
        control_supplier_key = supplier_key or items[0]["supplier_key"]
        return {"summary": summary, "items": items, "control": self.store.order_control(str(oc), control_supplier_key)}

    def _can_send(self, item: dict[str, Any], cooldown_hours: int) -> bool:
        last = self.store.last_success(item["item_key"])
        if not last:
            return True
        if last.get("status") == "uncertain":
            return False
        if last["urgency"] != item["urgency"]:
            return True
        then = datetime.fromisoformat(last["sent_at"])
        return (datetime.now() - then).total_seconds() >= cooldown_hours * 3600

    def _message(self, supplier: dict[str, Any], items: list[dict[str, Any]], settings: dict[str, Any]) -> str:
        contact = supplier.get("contact_name") or ""
        greeting = f"Olá, {contact}. Tudo bem?" if contact else "Olá, tudo bem?"
        lines = [greeting, "", f"Aqui é {settings['sender_name']}, do setor de Compras.",
                 "Estamos acompanhando os pedidos abaixo e precisamos da confirmação da entrega:", ""]
        by_order: dict[tuple[str, str], list[dict[str, Any]]] = {}
        for item in items:
            by_order.setdefault((item["company"], item["oc"]), []).append(item)
        for (company, oc), order_items in by_order.items():
            highest = max(order_items, key=lambda item: URGENCY_RANK[item["urgency"]])
            due = datetime.fromisoformat(highest["due_date"]).strftime("%d/%m/%Y")
            if highest["days"] == 0:
                timing = "vence hoje"
            elif highest["days"] and highest["days"] > 0:
                timing = f"vence em {highest['days']} dia(s)"
            else:
                timing = f"atrasado há {abs(highest['days'] or 0)} dia(s)"
            lines.append(f"• OC {oc} — {company}")
            lines.append(f"  Previsão: {due} ({timing})")
            for item in order_items[:8]:
                quantity = f" — Qtd. {item['quantity']:g}" if item.get("quantity") is not None else ""
                lines.append(f"  - {item['description']}{quantity}")
            if len(order_items) > 8:
                lines.append(f"  - e mais {len(order_items) - 8} item(ns)")
            lines.append("")
        lines.extend(["Por favor, confirme a data atualizada de entrega e informe qualquer impedimento.", "", settings["message_signature"]])
        return "\n".join(lines).strip()

    def groups(self, include_blocked: bool = False, buyer: str = "") -> list[dict[str, Any]]:
        settings = self.store.settings()
        grouped: dict[str, dict[str, Any]] = {}
        buyer_filter = str(buyer or "").strip()
        order_items = self.orders()
        last_successes = self.store.last_successes([item["item_key"] for item in order_items])
        for item in order_items:
            if buyer_filter and not person_matches(item.get("buyer"), buyer_filter):
                continue
            bpm_code = item.get("order_bpm_status_code")
            bpm_label = normalize(item.get("order_bpm_status", ""))
            if bpm_code is not None and int(bpm_code) != 3:
                continue
            if bpm_code is None and (
                any(word in bpm_label for word in ("RECUS", "REPROV", "REJEIT"))
                or any(word in bpm_label for word in ("AGUARD", "PEND", "EM APROVAC", "NAO APROV"))
            ):
                continue
            if not item["eligible"] or not item.get("active", 1):
                continue
            blocked_reason = ""
            if not item.get("phone"):
                blocked_reason = "Telefone não cadastrado"
            else:
                last = last_successes.get(item["item_key"])
                if last:
                    if last.get("status") == "uncertain":
                        blocked_reason = "Aguardando intervalo anti-spam"
                    elif last["urgency"] == item["urgency"]:
                        then = datetime.fromisoformat(last["sent_at"])
                        if (datetime.now() - then).total_seconds() < settings["cooldown_hours"] * 3600:
                            blocked_reason = "Aguardando intervalo anti-spam"
            if blocked_reason and not include_blocked:
                continue
            group = grouped.setdefault(item["supplier_key"], {
                "supplier_key": item["supplier_key"], "supplier_name": item["supplier_name"],
                "phone": item.get("phone") or "", "contact_name": item.get("contact_name") or "",
                "items": [], "blocked_reasons": set()
            })
            item["_blocked_reason"] = blocked_reason
            group["items"].append(item)
            if blocked_reason:
                group["blocked_reasons"].add(blocked_reason)
        result = []
        for group in grouped.values():
            ready_items = [item for item in group["items"] if not item.get("_blocked_reason")]
            message_items = ready_items or group["items"]
            group["items"] = message_items
            group["urgency"] = max((item["urgency"] for item in message_items), key=lambda value: URGENCY_RANK[value])
            group["message"] = self._message(group, message_items, settings)
            group["order_count"] = len({item["oc"] for item in message_items})
            group["item_count"] = len(message_items)
            group["blocked_reasons"] = [] if ready_items else sorted(group["blocked_reasons"])
            result.append(group)
        return sorted(result, key=lambda group: (-URGENCY_RANK[group["urgency"]], group["supplier_name"]))

    def dashboard(self, buyer: str = "") -> dict[str, Any]:
        buyer = str(buyer or "").strip()
        orders = [item for item in self.orders() if not buyer or person_matches(item.get("buyer"), buyer)]
        summaries = self.order_summaries(buyer=buyer)
        open_summaries = [order for order in summaries if order["urgency"] != "completed"]
        counts = {key: 0 for key in URGENCY_LABEL}
        for order in summaries:
            counts[order["urgency"]] += 1
        all_groups = self.groups(include_blocked=True, buyer=buyer)
        ready = [group for group in all_groups if not group["blocked_reasons"]]
        current_supplier_keys = {item["supplier_key"] for item in orders if item["urgency"] != "completed"}
        suppliers = {supplier["supplier_key"]: supplier for supplier in self.store.suppliers()}
        settings = self.store.settings()
        return {
            "counts": counts,
            "total_items": len(orders),
            "total_orders": len(summaries),
            "open_orders": len(open_summaries),
            "partial_items": sum(1 for item in orders if self._item_attendance(item) == "partial"),
            "total_value_open": sum(item.get("value_total") or 0 for item in orders if item["urgency"] != "completed"),
            "buyer_filter": buyer,
            "critical_after_days": settings["critical_after_days"],
            "ready_messages": len(ready),
            "suppliers_without_phone": sum(1 for key in current_supplier_keys if not suppliers.get(key, {}).get("phone")),
            "control_counts": {
                "blank": sum(1 for order in open_summaries if not order.get("control_status")),
                "sent": sum(1 for order in open_summaries if order.get("control_status") == "sent"),
                "card_payment": sum(1 for order in open_summaries if order.get("control_status") == "card_payment"),
                "pending": sum(1 for order in open_summaries if order.get("control_status") == "pending"),
                "request_collection": sum(1 for order in open_summaries if order.get("control_status") == "request_collection"),
                "awaiting_collection": sum(1 for order in open_summaries if order.get("control_status") == "awaiting_collection"),
            },
            "priority_orders": open_summaries[:12],
            "dashboard_orders": open_summaries[:1000],
            "last_import": self.store.latest_import(),
            "last_history": self.store.recent_history(8),
        }

    def send(self, supplier_keys: list[str] | None = None, force_simulation: bool | None = None, message_overrides: dict[str, str] | None = None, progress_callback=None, job_id: str | None = None, buyer: str = "") -> dict[str, Any]:
        with self.job_lock:
            active_job = self.job_state.get("job_id") if self.job_state.get("active") else None
        if active_job and job_id != active_job:
            raise ValueError("Já existe um lote em execução ou aguardando conexão.")
        if not self.send_lock.acquire(blocking=False):
            raise ValueError("Já existe um lote em execução ou aguardando conexão.")
        try:
            return self._send(supplier_keys, force_simulation, message_overrides, progress_callback, buyer=buyer)
        finally:
            self.send_lock.release()

    def _send(self, supplier_keys=None, force_simulation=None, message_overrides=None, progress_callback=None, buyer: str = ""):
        settings = self.store.settings()
        simulation = settings["simulation"] if force_simulation is None else bool(force_simulation)
        if supplier_keys is None:
            selected = None
        else:
            if not isinstance(supplier_keys, list) or len(supplier_keys) > 500:
                raise ValueError("Seleção de fornecedores inválida.")
            cleaned = [str(key).strip() for key in supplier_keys if str(key).strip()]
            if not cleaned:
                raise ValueError("Nenhum fornecedor selecionado. O envio foi cancelado por segurança.")
            selected = set(cleaned)

        overrides: dict[str, str] = {}
        if message_overrides is not None:
            if not isinstance(message_overrides, dict) or len(message_overrides) > 200:
                raise ValueError("Edições de mensagem inválidas.")
            for raw_key, raw_message in message_overrides.items():
                key = str(raw_key).strip()
                if not key or len(key) > 256 or not isinstance(raw_message, str):
                    raise ValueError("Edição de mensagem inválida.")
                message = raw_message.strip()
                if not message:
                    raise ValueError("A mensagem editada não pode ficar vazia.")
                if len(message) > 60000:
                    raise ValueError("A mensagem editada é muito grande.")
                overrides[key] = message

        def apply_overrides(groups_to_update):
            for group in groups_to_update:
                edited = overrides.get(group["supplier_key"])
                if edited is not None:
                    group["message"] = edited
            return groups_to_update

        groups = [group for group in self._groups_for_buyer(buyer) if selected is None or group["supplier_key"] in selected]
        if groups and not simulation:
            whatsapp_request("/wait")
            # Re-check eligibility after the user connects; never send an outdated snapshot.
            groups = [group for group in self._groups_for_buyer(buyer) if selected is None or group["supplier_key"] in selected]
        groups = apply_overrides(groups)
        batch_id = str(uuid.uuid4())
        sent = failed = simulated = 0
        results = []
        aborted_error = None
        ensure_ready_before_next = False
        if progress_callback:
            progress_callback({
                "phase": "sending", "total": len(groups), "processed": 0,
                "sent": 0, "failed": 0, "uncertain": 0, "simulated": 0,
                "current_supplier": None,
            })
        for index, group in enumerate(groups):
            if progress_callback:
                progress_callback({"current_supplier": group["supplier_name"], "phase": "sending"})
            # Antes de cada fornecedor, confirme silenciosamente que a sessão segue pronta.
            # Isso permite que o monitor restaure uma queda entre duas mensagens sem
            # exigir que o usuário reinicie manualmente o lote.
            if not simulation and index > 0 and ensure_ready_before_next:
                try:
                    whatsapp_request("/wait")
                    ensure_ready_before_next = False
                except Exception as exc:
                    aborted_error = str(exc)
                    break
            status, error = "simulated", None
            followup_id = None
            if simulation:
                simulated += 1
            else:
                # A crash or lost response must not cause an automatic duplicate.
                followup_id = self.store.record_followup(batch_id, group, "uncertain", "Envio em andamento ou interrompido. Confira a conversa antes de reenviar.")
                try:
                    result = whatsapp_request("/send", {"phone": group["phone"], "message": group["message"]})
                    status = result.get("status", "uncertain")
                    if status not in {"sent", "failed", "uncertain"}:
                        status = "uncertain"
                    error = result.get("error")
                    if status == "sent":
                        resolved_phone = result.get("resolved_phone")
                        if resolved_phone and resolved_phone != group.get("phone"):
                            self.store.save_supplier({
                                "supplier_key": group["supplier_key"],
                                "display_name": group["supplier_name"],
                                "phone": resolved_phone,
                            })
                            group["phone"] = resolved_phone
                        sent += 1
                        time.sleep(2)
                    else:
                        failed += 1
                except Exception as exc:
                    status, error = "uncertain", "Resposta interrompida. Confira a conversa no WhatsApp antes de liberar outro envio."
                    failed += 1
            if followup_id is None:
                followup_id = self.store.record_followup(batch_id, group, status, error)
            else:
                self.store.finish_followup(followup_id, status, error)
            results.append({"id": followup_id, "supplier_name": group["supplier_name"], "status": status, "error": error})
            self.store.log_event("followup_result", {
                "batch_id": batch_id, "supplier_key": group["supplier_key"],
                "status": status, "error": error, "simulation": simulation
            }, "error" if status in {"failed", "uncertain"} else "info")
            uncertain_count = sum(1 for item in results if item["status"] == "uncertain")
            error_text = normalize(error or "")
            ensure_ready_before_next = (
                not simulation and (status == "uncertain" or "CONEX" in error_text or "WHATSAPP NAO CONECT" in error_text)
            )
            if progress_callback:
                progress_callback({
                    "processed": len(results), "sent": sent, "failed": failed,
                    "uncertain": uncertain_count, "simulated": simulated,
                })
            # Falhas comprovadamente anteriores ao envio e envios incertos ficam
            # isolados no fornecedor atual. O lote segue para os demais sem repetir
            # a mensagem incerta. O intervalo anti-spam impede nova tentativa automática.
        return {
            "batch_id": batch_id, "sent": sent, "simulated": simulated, "failed": failed,
            "processed": len(results), "selected": len(groups),
            "uncertain": sum(1 for item in results if item["status"] == "uncertain"),
            "aborted": bool(aborted_error), "aborted_error": aborted_error,
            "results": results
        }


class AutoScheduler(threading.Thread):
    def __init__(self, store: Store, importer: WorkbookImporter, service: FollowUpService, poll_seconds: int = 30):
        super().__init__(name="followup-auto-scheduler", daemon=True)
        self.store = store
        self.importer = importer
        self.service = service
        self.poll_seconds = poll_seconds
        self.stop_event = threading.Event()

    @staticmethod
    def should_run(settings: dict[str, Any], now: datetime) -> bool:
        if not (
            settings.get("automatic_enabled")
            and not settings.get("simulation")
            and settings.get("last_auto_run") != now.date().isoformat()
            and now.strftime("%H:%M") >= settings.get("schedule_time", "09:00")
        ):
            return False
        last_attempt = settings.get("last_auto_attempt", "")
        if last_attempt:
            try:
                attempt = datetime.fromisoformat(last_attempt)
                if attempt.date() == now.date() and (now - attempt).total_seconds() < 3600:
                    return False
            except ValueError:
                pass
        return True

    def run_once_if_due(self, now: datetime | None = None) -> bool:
        current = now or datetime.now()
        settings = self.store.settings()
        if not self.should_run(settings, current):
            return False
        try:
            self.store.set_internal("last_auto_attempt", current.isoformat(timespec="seconds"))
            workbook_path = settings.get("last_workbook_path", "")
            if not workbook_path:
                raise RuntimeError("Nenhuma planilha foi importada para a execução automática.")
            self.importer.import_file(workbook_path)
            sender = getattr(self.service, "send_persistent_and_wait", self.service.send)
            result = sender(force_simulation=False, buyer=settings.get("buyer_filter", ""))
            print(json.dumps({"event": "automatic_run", "result": result}, ensure_ascii=False), flush=True)
            self.store.log_event("automatic_run", result, "error" if result.get("failed") or result.get("uncertain") else "info")
            # Mark the day complete only after a clean run. Failures can be fixed
            # and retried later; uncertain sends require human review first.
            if not result.get("failed") and not result.get("uncertain"):
                self.store.set_internal("last_auto_run", current.date().isoformat())
        except Exception as exc:
            self.store.log_event("automatic_error", {"error": str(exc)}, "error")
            print(json.dumps({"event": "automatic_error", "error": str(exc)}, ensure_ascii=False), file=sys.stderr, flush=True)
        return True

    def run(self):
        while not self.stop_event.is_set():
            self.run_once_if_due()
            self.stop_event.wait(self.poll_seconds)


class WhatsAppSupervisor(threading.Thread):
    """Silent backend heartbeat for the local WhatsApp bridge.

    The bridge owns reconnection and session state. This supervisor only nudges
    its health check periodically. It never surfaces popups or changes business
    rules, and a user pause is respected by the bridge.
    """
    def __init__(self, poll_seconds: int = 20):
        super().__init__(name="whatsapp-backend-supervisor", daemon=True)
        self.poll_seconds = poll_seconds
        self.stop_event = threading.Event()

    def run(self):
        # Give Electron a few seconds to finish the normal application startup.
        self.stop_event.wait(5)
        while not self.stop_event.is_set():
            try:
                whatsapp_request("/health")
            except Exception:
                # The Node-side connection manager owns recovery. A temporary
                # bridge failure must stay silent and be tried again later.
                pass
            self.stop_event.wait(self.poll_seconds)


class ApiHandler(BaseHTTPRequestHandler):
    store: Store
    importer: WorkbookImporter
    service: FollowUpService
    token: str

    def log_message(self, _format, *_args):
        return

    def _json(self, status: int, payload: Any):
        data = json.dumps(payload, ensure_ascii=False, default=str).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(data)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(data)

    def _authorized(self) -> bool:
        return bool(self.token) and self.headers.get("X-FollowUp-Token") == self.token

    def _body(self) -> dict[str, Any]:
        length = int(self.headers.get("Content-Length", "0"))
        return json.loads(self.rfile.read(length) or b"{}")

    def do_GET(self):
        if not self._authorized():
            return self._json(401, {"error": "Não autorizado."})
        parsed = urlparse(self.path)
        try:
            if parsed.path == "/health":
                return self._json(200, {"ok": True, "app": APP_NAME})
            if parsed.path == "/dashboard":
                query = parse_qs(parsed.query)
                return self._json(200, self.service.dashboard(query.get("buyer", [""])[0]))
            if parsed.path == "/orders":
                query = parse_qs(parsed.query)
                return self._json(200, {"orders": self.service.order_summaries(
                    buyer=query.get("buyer", [""])[0], company=query.get("company", [""])[0],
                    urgency=query.get("urgency", [""])[0], search=query.get("search", [""])[0],
                    control_status=query.get("control_status", ["all"])[0],
                    attendance_status=query.get("attendance_status", ["all"])[0],
                )})
            if parsed.path == "/order":
                query = parse_qs(parsed.query)
                return self._json(200, self.service.order_detail(
                    query.get("oc", [""])[0], query.get("supplier_key", [""])[0]
                ))
            if parsed.path == "/filters":
                return self._json(200, self.service.filters())
            if parsed.path == "/suppliers":
                return self._json(200, {"suppliers": self.store.suppliers()})
            if parsed.path == "/preview":
                query = parse_qs(parsed.query)
                return self._json(200, {"groups": self.service.groups(include_blocked=True, buyer=query.get("buyer", [""])[0])})
            if parsed.path == "/history":
                return self._json(200, {"history": self.store.recent_history(200)})
            if parsed.path == "/settings":
                return self._json(200, self.store.settings())
            if parsed.path == "/send-status":
                query = parse_qs(parsed.query)
                return self._json(200, self.service.send_status(query.get("job_id", [""])[0] or None))
            return self._json(404, {"error": "Rota não encontrada."})
        except Exception as exc:
            return self._json(500, {"error": str(exc)})

    def do_POST(self):
        if not self._authorized():
            return self._json(401, {"error": "Não autorizado."})
        try:
            if self.path == "/prepare-update":
                if self.service.send_status().get("active") or not self.service.send_lock.acquire(blocking=False):
                    return self._json(409, {"error": "Envio em andamento."})
                # Retain the lock until this process exits: scheduler cannot start a new send.
                return self._json(200, {"ok": True})
            body = self._body()
            if self.path == "/import":
                return self._json(200, self.importer.import_file(body.get("path", "")))
            if self.path == "/supplier":
                return self._json(200, self.store.save_supplier(body))
            if self.path == "/order-control":
                return self._json(200, self.store.save_order_control(body))
            if self.path == "/settings":
                return self._json(200, self.store.save_settings(body))
            if self.path == "/followup-reviewed":
                if self.service.send_lock.locked():
                    raise ValueError("Aguarde o lote terminar antes de revisar o envio.")
                return self._json(200, self.store.review_followup(body.get("id")))
            if self.path == "/send-start":
                return self._json(202, self.service.start_send(
                    body.get("supplier_keys"), body.get("simulation"), body.get("message_overrides"), body.get("buyer", "")
                ))
            if self.path == "/send":
                return self._json(200, self.service.send(
                    body.get("supplier_keys"), body.get("simulation"), body.get("message_overrides"), buyer=body.get("buyer", "")
                ))
            return self._json(404, {"error": "Rota não encontrada."})
        except ValueError as exc:
            return self._json(400, {"error": str(exc)})
        except Exception as exc:
            return self._json(500, {"error": str(exc)})


def serve(port: int):
    data_dir = Path(os.environ.get("FOLLOWUP_DATA_DIR", Path.home() / ".vyzium"))
    store = Store(data_dir / "followup.db")
    ApiHandler.store = store
    ApiHandler.importer = WorkbookImporter(store)
    ApiHandler.service = FollowUpService(store)
    ApiHandler.token = os.environ.get("FOLLOWUP_API_TOKEN", "")
    if not ApiHandler.token:
        raise RuntimeError("Token local ausente.")
    scheduler = AutoScheduler(store, ApiHandler.importer, ApiHandler.service)
    scheduler.start()
    whatsapp_supervisor = WhatsAppSupervisor()
    whatsapp_supervisor.start()
    server = ThreadingHTTPServer(("127.0.0.1", port), ApiHandler)
    print(json.dumps({"event": "ready", "port": server.server_address[1]}), flush=True)
    server.serve_forever()


def main():
    parser = argparse.ArgumentParser(description=APP_NAME)
    subparsers = parser.add_subparsers(dest="command", required=True)
    serve_parser = subparsers.add_parser("serve")
    serve_parser.add_argument("--port", type=int, default=0)
    args = parser.parse_args()
    if args.command == "serve":
        serve(args.port)


if __name__ == "__main__":
    main()
