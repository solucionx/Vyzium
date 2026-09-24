from __future__ import annotations

import argparse
import base64
import csv
import hashlib
import hmac
import io
import json
import os
import re
import sqlite3
import sys
import threading
import unicodedata
import uuid
from contextlib import contextmanager
from datetime import date, datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse, parse_qs
from urllib.request import Request, urlopen
from urllib.error import HTTPError, URLError
from openpyxl import load_workbook
from workbook_formats import open_book, REPORT_REQUIRED, adapt_report, approval_deadline, normalize_header
from data_safety import DataIntegrityError, DataSafetyManager
from secure_sqlite import connect as secure_connect, key_from_env


APP_VERSION = os.environ.get('VYZIUM_APP_VERSION', '3.2.6')
DB_SCHEMA_VERSION = 1

def norm(value):
    return normalize_header(value)


# CNPJ used only in supplier-facing quotation messages. The SCI remains an
# internal Vyzium identifier and must never be exposed in those messages.
HOTEL_CNPJ = {
    'MAGNA PRAIA': '02.333.096/0001-35',
    'MAGNA PRAIA HOTEL': '02.333.096/0001-35',
    'CARMEL TAIBA': '27.708.448/0001-10',
    'CARMEL TAIBA EXCLUSIVE RESORT HOTEIS LTDA': '27.708.448/0001-10',
    'CHARME HOSPEDAGEM': '27.794.852/0001-54',
    'CARMEL RESORT HOSPEDAGEM LTDA EPP': '27.794.852/0001-54',
    'CARMEL CUMBUCO': '19.253.187/0001-63',
    'CARMEL WIND RESORT LTDA CUMBUCO': '19.253.187/0001-63',
    'CM SERVICOS': '35.428.047/0001-35',
    'CM CENTRAL DE SERVICOS ADMINISTRATIVOS LTDA': '35.428.047/0001-35',
    'CARMEL ICARAIZINHO': '45.862.118/0001-67',
    'CARMEL ICARAIZINHO RESORT LTDA': '45.862.118/0001-67',
}

AWARD_REASONS = (
    'Prazo de entrega',
    'Disponibilidade imediata',
    'Qualidade',
    'Frete',
    'Condição de pagamento',
    'Histórico do fornecedor',
    'Necessidade do hotel',
    'Outro',
)


def format_brl(value):
    value = money(Decimal(value))
    integer, decimal = f"{value:.2f}".split('.')
    chunks = []
    while integer:
        chunks.append(integer[-3:])
        integer = integer[:-3]
    return 'R$ ' + '.'.join(reversed(chunks)) + ',' + decimal



def hotel_cnpj(company):
    return HOTEL_CNPJ.get(norm(company), '')


def text(value):
    if value is None:
        return ''
    return str(int(value)) if isinstance(value, float) and value.is_integer() else str(value).strip()


def number(value):
    if isinstance(value, bool):
        raise ValueError('Valor numérico inválido.')
    raw = str(value).strip().replace('R$', '').replace(' ', '')
    if ',' in raw:
        raw = raw.replace('.', '').replace(',', '.')
    try:
        result = Decimal(raw)
    except InvalidOperation:
        raise ValueError('Informe um número válido.')
    if not result.is_finite() or result < 0 or result > Decimal('1000000000000'):
        raise ValueError('Valor fora do intervalo permitido.')
    return result


def money(value):
    return value.quantize(Decimal('.01'), rounding=ROUND_HALF_UP)


def iso(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%Y-%m-%d')
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(str(value)[:10], fmt).date().isoformat()
        except ValueError:
            pass
    return ''


def now():
    return datetime.now().isoformat(timespec='seconds')


def load_items(path):
    """Read the raw SCI export, never treating receipt rows as new demand."""
    book = open_book(path, load_workbook)
    groups = {}
    rows_seen = 0
    try:
        source = None
        headers = []
        report_format = False
        required = {'FKEMPRESA', 'EMPRESA', 'IDSCI', 'IDITEMDASCI', 'DESCRICAOARTIGO',
                    'QUANTIDADESCI', 'UNIDADEDEMEDIDASCI', 'IDORDEMDECOMPRA',
                    'NMSTATUSDOITEMDASCI', 'NMSTATUSBPMSCI', 'COMPRADOR'}
        recognized = required | REPORT_REQUIRED
        for sheet in book:
            sheet.reset_dimensions()  # SCI export declares A1:A1 despite containing all rows.
            best = None
            for row_index, values in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True), 1):
                candidate = [norm(v) for v in values]
                candidate_set = {h for h in candidate if h}
                candidate_report = REPORT_REQUIRED.issubset(candidate_set)
                candidate_raw = required.issubset(candidate_set)
                score = len(candidate_set & recognized) + (100 if candidate_report or candidate_raw else 0)
                if best is None or score > best[0]:
                    best = (score, row_index, candidate, candidate_report, candidate_raw)
            if not best or not (best[3] or best[4]):
                continue
            _, header_row, headers, report_format, _ = best
            duplicates = sorted({h for h in headers if h and headers.count(h) > 1 and h in recognized})
            if duplicates:
                raise ValueError('Cabeçalhos duplicados/ambíguos: ' + ', '.join(duplicates) + '. A base anterior foi preservada.')
            source = sheet.iter_rows(min_row=header_row + 1, values_only=True)
            break
        if source is None:
            raise ValueError('Cabeçalhos não reconhecidos nas primeiras 30 linhas. Use o relatório com Número da SCI, Comprador SCI, Status da SCI e Status BPM SCI, ou a BASE SCI original.')
        for row in source:
            if not any(v is not None for v in row):
                continue
            # Some exports repeat the header inside the data region (and custom
            # spreadsheet adapters may ignore min_row). Never count a header as
            # a purchase item.
            if [norm(v) for v in row] == headers:
                continue
            r = dict(zip(headers, row))
            if report_format:
                r = adapt_report(r, norm, text)
                if r is None:
                    continue
            rows_seen += 1
            key = '|'.join(text(r.get(k)) for k in ('FKEMPRESA', 'IDSCI', 'IDITEMDASCI'))
            if not text(r.get('IDITEMDASCI')) or not text(r.get('IDSCI')):
                raise ValueError(f'Linha {rows_seen + 1}: SCI ou identificação do item ausente.')
            groups.setdefault(key, []).append(r)
    finally:
        book.close()
    result = []
    deadline_conflicts = 0
    excluded = {'with_order': 0, 'closed_or_rejected': 0, 'other_status': 0, 'invalid_quantity': 0, 'conflict': 0}
    for key, rows in groups.items():
        if len(rows) > 1 and rows[0].get('REPORT_IDENTITY'):
            excluded['conflict'] += 1
            continue
        if any(text(r.get('IDORDEMDECOMPRA')) not in ('', '0', '0.0') for r in rows):
            excluded['with_order'] += 1
            continue
        if any(text(r.get('NMSTATUSDOITEMDASCI')) in ('4', '6') or text(r.get('NMSTATUSBPMSCI')) in ('2', '4') for r in rows):
            excluded['closed_or_rejected'] += 1
            continue
        r = rows[0]
        if any(text(x.get('NMSTATUSDOITEMDASCI')) not in ('0', '2') for x in rows):
            excluded['other_status'] += 1
            continue
        comparable = ('QUANTIDADESCI', 'UNIDADEDEMEDIDASCI', 'DESCRICAOARTIGO', 'COMPRADOR', 'NMSTATUSDOITEMDASCI', 'NMSTATUSBPMSCI')
        if any(any(text(x.get(f)) != text(r.get(f)) for f in comparable) for x in rows[1:]):
            excluded['conflict'] += 1
            continue
        try:
            qty = number(r.get('QUANTIDADESCI'))
            if qty <= 0:
                raise ValueError()
        except ValueError:
            excluded['invalid_quantity'] += 1
            continue
        approved_at, due, mismatch = approval_deadline(r, iso)
        deadline_conflicts += int(mismatch)
        result.append({'id': key, 'company': text(r.get('EMPRESA')), 'sci': text(r.get('IDSCI')),
                       'article': text(r.get('CODIGOARTIGO')), 'description': text(r.get('DESCRICAOARTIGO')),
                       'quantity': str(qty), 'unit': text(r.get('UNIDADEDEMEDIDASCI')),
                       'buyer': text(r.get('COMPRADOR')), 'group': text(r.get('DESCRICAOGRUPO')),
                       'needed': due, 'approved_at': approved_at, 'deadline_rule': 'approval_12_days',
                       'note': text(r.get('NOTE')), 'purchase_type': text(r.get('PURCHASE_TYPE')),
                       'issued': iso(r.get('DATAEMISSAOSCI')),
                       'urgent': norm(r.get('URGENTE')) in ('TRUE', '1', 'SIM', 'URGENTE'),
                       'status': 'pending' if text(r.get('NMSTATUSDOITEMDASCI')) == '0' else 'quoting',
                       'approval': 'approved' if text(r.get('NMSTATUSBPMSCI')) == '3' else 'waiting'})
    return result, {'rows': rows_seen, 'unique_items': len(groups), 'eligible': len(result), 'excluded': excluded,
                    'format': 'approval_report' if report_format else 'raw_sci', 'deadline_conflicts': deadline_conflicts,
                    'missing_deadline': sum(not i['needed'] for i in result)}


def _saving_target(data):
    raw = data.get('saving_target', '5')
    target = number('5' if raw in ('', None) else raw)
    if target > Decimal('99.99'):
        raise ValueError('A meta de saving deve ficar entre 0% e 99,99%.')
    return target


def evaluate(data):
    """Evaluate prices without conflating the financial best with the buyer's final choice.

    The cheapest valid final quote remains the financial reference. A buyer may explicitly
    choose another quoted supplier for an operational reason (delivery, availability, etc.).
    Saving is always measured against the initial price of the supplier that was actually
    chosen, while ``opportunity_cost`` records any premium versus the lowest final quote.
    """
    lines = []
    totals = {s['id']: {'name': s['name'], 'won': 0, 'net': Decimal(0), 'saving': Decimal(0)} for s in data['suppliers']}
    target_percent = _saving_target(data)
    total_opportunity_cost = Decimal(0)
    total_initial_selected = Decimal(0)
    target_map_value = Decimal(0)
    target_lines = 0
    awards = data.get('awards') or {}

    for item in data['items']:
        quotes = []
        for supplier in data['suppliers']:
            q = data['quotes'].get(item['id'], {}).get(supplier['id'], {})
            if q.get('price') in (None, ''):
                continue
            price = number(q['price'])
            if price <= 0:
                raise ValueError('Preço precisa ser maior que zero; deixe vazio quando não houver cotação.')
            gross = money(price * number(item['quantity']))
            if 'negotiated' in q:
                final_price = price if q['negotiated'] in ('', None) else number(q['negotiated'])
                if final_price <= 0 or final_price > price:
                    raise ValueError('O preço negociado deve ser maior que zero e não pode superar o preço inicial.')
                net = money(final_price * number(item['quantity']))
                saving = gross - net
                percent = money((price - final_price) / price * 100)
            else:
                # Preserve old maps: their discount fields still mean percent or a line amount.
                discount = number(q.get('discount') or '0')
                kind = q.get('kind', 'percent')
                if kind not in ('percent', 'amount'):
                    raise ValueError('Tipo de desconto inválido.')
                if (kind == 'percent' and discount > 100) or (kind == 'amount' and discount > gross):
                    raise ValueError('Desconto não pode ultrapassar o valor do item.')
                saving = money(gross * discount / 100) if kind == 'percent' else money(discount)
                net = gross - saving
                final_price = net / number(item['quantity'])
                percent = money(saving / gross * 100) if gross else Decimal(0)
            quotes.append({'supplier_id': supplier['id'], 'supplier': supplier['name'], 'gross': str(gross),
                           'saving': str(saving), 'net': str(net), 'initial_price': str(price),
                           'final_price': str(final_price), 'discount_percent': str(percent),
                           'delivery': text(q.get('delivery'))[:120]})

        best = min((Decimal(q['net']) for q in quotes), default=None)
        winners = [q for q in quotes if Decimal(q['net']) == best]
        best_initial = min((Decimal(q['initial_price']) for q in quotes), default=None)
        target_price = None
        if best_initial is not None:
            target_price = money(best_initial * (Decimal('1') - target_percent / Decimal('100')))
            target_map_value += money(target_price * number(item['quantity']))
            target_lines += 1
            for quote in quotes:
                current = Decimal(quote['final_price'])
                remaining = max(Decimal(0), current - target_price)
                quote['target_price'] = str(target_price)
                quote['target_reduction_unit'] = str(money(remaining))
                quote['required_reduction_percent'] = str(money(remaining / current * 100) if current else Decimal(0))
                quote['target_met'] = current <= target_price

        award = awards.get(item['id']) if isinstance(awards, dict) else None
        chosen = None
        manual_selection = False
        selection_reason = ''
        selection_note = ''
        if isinstance(award, dict) and award.get('supplier_id'):
            chosen = next((q for q in quotes if q['supplier_id'] == award.get('supplier_id')), None)
            manual_selection = chosen is not None
            selection_reason = text(award.get('reason'))[:120]
            selection_note = text(award.get('note'))[:500]
        if chosen is None:
            selected = data.get('choices', {}).get(item['id'])
            chosen = next((q for q in winners if q['supplier_id'] == selected), None)
            if len(winners) == 1:
                chosen = winners[0]

        opportunity_cost = Decimal(0)
        if chosen:
            total = totals[chosen['supplier_id']]
            total['won'] += 1
            total['net'] += Decimal(chosen['net'])
            total['saving'] += Decimal(chosen['saving'])
            total_initial_selected += Decimal(chosen['gross'])
            if best is not None:
                opportunity_cost = max(Decimal(0), Decimal(chosen['net']) - best)
                total_opportunity_cost += opportunity_cost

        lines.append({'id': item['id'], 'quotes': quotes, 'winners': winners, 'financial_winners': winners,
                      'chosen': chosen, 'manual_selection': manual_selection,
                      'selection_reason': selection_reason, 'selection_note': selection_note,
                      'opportunity_cost': str(money(opportunity_cost)),
                      'target_price': str(target_price) if target_price is not None else '',
                      'state': 'unquoted' if not winners else ('tie' if not chosen else 'winner')})

    net_total = sum((t['net'] for t in totals.values()), Decimal(0))
    saving_total = sum((t['saving'] for t in totals.values()), Decimal(0))
    saving_percent = money(saving_total / total_initial_selected * 100) if total_initial_selected else Decimal(0)
    target_saving_amount = money(total_initial_selected * target_percent / 100) if total_initial_selected else Decimal(0)
    target_gap = max(Decimal(0), target_saving_amount - saving_total)
    return {'lines': lines,
            'suppliers': [{**t, 'id': sid, 'net': str(t['net']), 'saving': str(t['saving'])} for sid, t in totals.items()],
            'net': str(net_total), 'saving': str(saving_total), 'gross': str(total_initial_selected),
            'saving_percent': str(saving_percent), 'saving_target_percent': str(target_percent),
            'saving_target_amount': str(target_saving_amount), 'saving_target_gap': str(money(target_gap)),
            'target_map_value': str(money(target_map_value)), 'target_lines': target_lines,
            'opportunity_cost': str(money(total_opportunity_cost)),
            'unquoted': sum(x['state'] == 'unquoted' for x in lines),
            'ties': sum(x['state'] == 'tie' for x in lines)}


def whatsapp_request(route, body=None):
    base = os.environ.get('FOLLOWUP_WHATSAPP_URL', '')
    if not re.fullmatch(r'http://127\.0\.0\.1:\d+', base):
        raise ValueError('Abra o aplicativo para conectar o WhatsApp.')
    req = Request(base + route, data=json.dumps(body or {}).encode(),
                  headers={'Content-Type': 'application/json', 'X-FollowUp-Token': os.environ.get('FOLLOWUP_API_TOKEN', '')})
    try:
        with urlopen(req, timeout=135 if route == '/wait' else 85) as response:
            return json.load(response)
    except HTTPError as exc:
        try:
            message = json.load(exc).get('error', 'Erro na conexão do WhatsApp.')
        except Exception:
            message = 'Erro na conexão do WhatsApp.'
        raise RuntimeError(message) from exc
    except (URLError, TimeoutError, ConnectionError, json.JSONDecodeError) as exc:
        raise RuntimeError('Não foi possível falar com a ponte do WhatsApp. Verifique a conexão e tente novamente.') from exc


class Store:
    def __init__(self, directory):
        directory = Path(directory).expanduser().resolve()
        directory.mkdir(parents=True, exist_ok=True)
        self.path = str(directory / 'compras.sqlite3')
        self.db_key_hex = key_from_env()
        self.safety = DataSafetyManager(self.path, 'compras', APP_VERSION, auto_retention=10, key_hex=self.db_key_hex)
        had_existing_database = self.safety.has_existing_database
        self.safety.assert_existing_integrity()
        self.lock = threading.RLock()
        self.send_lock = threading.Lock()
        if had_existing_database:
            # Snapshot through a read-only source before journal/schema maintenance.
            self.safety.ensure_version_backup()
            existing_schema = self.safety.existing_schema_version()
            if existing_schema > DB_SCHEMA_VERSION:
                raise DataIntegrityError(
                    f"Este banco usa o schema {existing_schema}, mais novo que o suportado por esta versão ({DB_SCHEMA_VERSION}). "
                    "O Vyzium bloqueou a abertura para evitar que uma versão antiga altere dados mais novos."
                )
        with self.db() as con:
            con.executescript('''CREATE TABLE IF NOT EXISTS items(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS maps(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS settings(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS messages(id TEXT PRIMARY KEY, data TEXT NOT NULL);
            CREATE TABLE IF NOT EXISTS schema_migrations(
                version INTEGER PRIMARY KEY, name TEXT NOT NULL, applied_at TEXT NOT NULL
            );''')
            con.execute('INSERT OR IGNORE INTO schema_migrations(version,name,applied_at) VALUES (?,?,?)',
                        (DB_SCHEMA_VERSION, 'data_safety_v1', now()))
            for mid, raw in con.execute('SELECT id,data FROM messages').fetchall():
                msg = json.loads(raw)
                if msg['status'] == 'sending':
                    msg.update(status='uncertain', error='Aplicativo encerrado durante envio. Confira a conversa.')
                    con.execute('UPDATE messages SET data=? WHERE id=?', (json.dumps(msg), mid))

    @contextmanager
    def db(self):
        con = secure_connect(self.path, key_hex=self.db_key_hex, timeout=30)
        try:
            con.execute('PRAGMA busy_timeout=10000')
            con.execute('PRAGMA foreign_keys=ON')
            con.execute('PRAGMA journal_mode=WAL')
            con.execute('PRAGMA synchronous=FULL')
            with con:
                yield con
        finally:
            con.close()

    def data_safety_status(self):
        with self.lock, self.db() as con:
            result = self.safety.status()
            row = con.execute('SELECT MAX(version) FROM schema_migrations').fetchone()
            result['schema_version'] = int((row or [0])[0] or 0)
            result['app_version'] = APP_VERSION
            return result

    def create_backup(self, reason='manual', automatic=False):
        with self.lock, self.db() as con:
            return self.safety.backup(reason=reason, source_connection=con, automatic=automatic)

    JSON_TABLES = frozenset({'items', 'maps', 'settings', 'messages'})

    @classmethod
    def _safe_table(cls, table):
        name = str(table or '')
        if name not in cls.JSON_TABLES:
            raise ValueError('Tabela interna inválida.')
        return name

    def all(self, table):
        table = self._safe_table(table)
        with self.db() as con:
            return [json.loads(r[0]) for r in con.execute(f'SELECT data FROM {table}')]

    def put(self, table, key, data):
        table = self._safe_table(table)
        with self.db() as con:
            con.execute(f'INSERT OR REPLACE INTO {table}(id,data) VALUES (?,?)', (key, json.dumps(data, ensure_ascii=False)))

    def get_map(self, mid):
        with self.db() as con:
            row = con.execute('SELECT data FROM maps WHERE id=?', (mid,)).fetchone()
        if not row:
            raise ValueError('Mapa não encontrado.')
        return json.loads(row[0])

    def settings(self):
        values = self.all('settings')
        return next((v for v in values if v.get('id') == 'ui'), {'id': 'ui', 'filters': {}})

    def import_file(self, path):
        if Path(path).suffix.lower() not in ('.xls', '.xlsx', '.xlsm'):
            raise ValueError('Escolha um arquivo XLS, XLSX ou XLSM.')
        items, report = load_items(path)
        with self.lock, self.db() as con:
            self.safety.backup(reason='pre-import', source_connection=con, automatic=True)
            con.execute('DELETE FROM items')
            con.executemany('INSERT INTO items VALUES (?,?)', [(i['id'], json.dumps(i, ensure_ascii=False)) for i in items])
            report.update(id='import', at=now(), filename=Path(path).name)
            con.execute('INSERT OR REPLACE INTO settings VALUES (?,?)', ('import', json.dumps(report)))
            check = [str(row[0]) for row in con.execute('PRAGMA quick_check').fetchall()]
            if check != ['ok']:
                raise DataIntegrityError('A nova importação não passou na verificação interna; alterações revertidas.')
        return report

    def catalog(self):
        items = self.all('items')
        assigned = {}
        for m in self.all('maps'):
            if m.get('archived'):
                continue
            for i in m['items']:
                assigned.setdefault(i['id'], []).append(m['name'])
        for i in items:
            i['maps'] = assigned.get(i['id'], [])
            if i.get('deadline_rule') != 'approval_12_days':
                i['needed'] = ''  # Old stored necessity dates must never masquerade as approval deadlines.
        return {'items': items, 'import': next((v for v in self.all('settings') if v.get('id') == 'import'), None)}

    def overview(self):
        catalog = self.catalog()
        items = catalog['items']
        today_value = date.today().isoformat()
        active_maps = [m for m in self.all('maps') if not m.get('archived')]
        scis = {(text(i.get('company')), text(i.get('sci'))) for i in items if text(i.get('sci'))}
        overdue = 0
        for item in items:
            needed = text(item.get('needed'))
            if needed and needed < today_value:
                overdue += 1
        return {
            'total_items': len(items),
            'total_scis': len(scis),
            'overdue_items': overdue,
            'mapped_items': sum(1 for i in items if i.get('maps')),
            'active_maps': len(active_maps),
            'last_import': catalog.get('import'),
        }

    def create_map(self, body):
        name = text(body.get('name'))[:160]
        ids = list(dict.fromkeys(body.get('ids', [])))
        if not name or not ids:
            raise ValueError('Informe um nome e selecione ao menos um item.')
        with self.lock:
            catalog = {i['id']: i for i in self.catalog()['items']}
            if any(i not in catalog for i in ids):
                raise ValueError('Um item deixou de estar disponível. Atualize a lista.')
            if any(catalog[i]['maps'] for i in ids):
                raise ValueError('Há item em outro mapa ativo. Conclua o mapa anterior antes de reutilizá-lo.')
            data = {'id': uuid.uuid4().hex, 'name': name, 'created': now(), 'revision': 1,
                    'items': [{**catalog[i], 'note': catalog[i].get('note', ''), 'purchase_type': catalog[i].get('purchase_type', '')} for i in ids],
                    'suppliers': [], 'quotes': {}, 'choices': {}, 'awards': {}, 'saving_target': '5', 'archived': False}
            self.put('maps', data['id'], data)
        return data

    def map_summaries(self):
        """Small searchable summaries for the map library without changing stored map data."""
        summaries = []
        for data in self.all('maps'):
            items = data.get('items', [])
            summaries.append({
                'id': data.get('id', ''),
                'name': data.get('name', ''),
                'created': data.get('created', ''),
                'updated': data.get('updated', ''),
                'completed_at': data.get('completed_at', ''),
                'count': len(items),
                # Keep the historical `archived` flag for backwards compatibility; in the UI it means concluded.
                'archived': bool(data.get('archived')),
                'scis': sorted({text(item.get('sci')) for item in items if text(item.get('sci'))}),
                'items': [text(item.get('description')) for item in items if text(item.get('description'))],
                'articles': [text(item.get('article')) for item in items if text(item.get('article'))],
            })
        summaries.sort(key=lambda m: (m.get('completed_at') or m.get('updated') or m.get('created') or '', m.get('created') or ''), reverse=True)
        return summaries

    def complete_map(self, mid):
        with self.lock:
            data = self.get_map(mid)
            if not data.get('archived'):
                stamp = now()
                data.update(archived=True, completed_at=stamp, updated=stamp, revision=data.get('revision', 0) + 1)
                self.put('maps', data['id'], data)
        return self.detail(mid)

    def delete_map(self, mid):
        """Permanently remove the map/quotes. Sent-message history is intentionally retained as an audit log."""
        with self.lock, self.db() as con:
            row = con.execute('SELECT data FROM maps WHERE id=?', (mid,)).fetchone()
            if not row:
                raise ValueError('Mapa não encontrado.')
            # Destructive action gets an automatic recovery point before deletion.
            self.safety.backup(reason='pre-delete-map', source_connection=con, automatic=True)
            con.execute('DELETE FROM maps WHERE id=?', (mid,))
            check = [str(row[0]) for row in con.execute('PRAGMA quick_check').fetchall()]
            if check != ['ok']:
                raise DataIntegrityError('A exclusão não passou na verificação interna; alterações revertidas.')
        return {'ok': True}

    def save_map(self, body):
        with self.lock:
            data = self.get_map(body.get('id'))
            if data['archived']:
                raise ValueError('Mapa concluído: somente consulta e exportação estão disponíveis.')
            if body.get('revision') != data['revision']:
                raise ValueError('O mapa foi alterado em outra tela. Abra novamente antes de salvar.')
            # Removal is explicit: a partial/older payload must never delete items by omission.
            remove_ids = body.get('remove_item_ids', [])
            if not isinstance(remove_ids, list) or any(not isinstance(i, str) for i in remove_ids):
                raise ValueError('Seleção de itens para remoção inválida.')
            removed = set(remove_ids)
            # Add only current, available catalog rows; never trust imported fields from the renderer.
            add_ids = body.get('add_item_ids', [])
            if (not isinstance(add_ids, list)
                    or any(not isinstance(i, str) or not i for i in add_ids)
                    or len(set(add_ids)) != len(add_ids)):
                raise ValueError('Seleção de itens para inclusão inválida.')
            added = set(add_ids)
            new_items = []
            if added:
                if removed:
                    raise ValueError('Adicione ou remova itens em operações separadas.')
                if self.send_lock.locked():
                    raise ValueError('Aguarde o envio terminar antes de adicionar itens.')
                if added.intersection(i['id'] for i in data['items']):
                    raise ValueError('Um item selecionado já está neste mapa. Atualize a seleção.')
                catalog = {i['id']: i for i in self.catalog()['items']}
                if any(i not in catalog for i in add_ids):
                    raise ValueError('Um item deixou de estar disponível na base. Feche e abra a seleção novamente.')
                if any(catalog[i]['maps'] for i in add_ids):
                    raise ValueError('Um item já está em outro mapa ativo. Remova-o de lá antes de adicionar aqui.')
                new_items = [{**catalog[i], 'note': catalog[i].get('note', ''),
                              'purchase_type': catalog[i].get('purchase_type', '')} for i in add_ids]
                # Inclusion may omit unchanged editable fields without erasing existing quotes/suppliers.
                body = {**data, **body}
            if removed:
                if self.send_lock.locked():
                    raise ValueError('Aguarde o envio terminar antes de remover um item.')
                if not removed.issubset({i['id'] for i in data['items']}):
                    raise ValueError('Um item não pertence mais a este mapa. Abra o mapa novamente.')
                data['items'] = [i for i in data['items'] if i['id'] not in removed]
                if not data['items']:
                    raise ValueError('O mapa precisa manter ao menos um item. Para remover o último, use Excluir mapa.')
            # Only user-owned fields are editable; imported identity and quantities remain authoritative.
            data['name'] = text(body.get('name', data['name']))[:160]
            if not data['name']:
                raise ValueError('Nome do mapa obrigatório.')

            target = number(body.get('saving_target', data.get('saving_target', '5')) or '5')
            if target > Decimal('99.99'):
                raise ValueError('A meta de saving deve ficar entre 0% e 99,99%.')
            data['saving_target'] = str(target)

            suppliers = []
            seen = set()
            names = set()
            for s in body.get('suppliers', []):
                sid, name = text(s.get('id')), text(s.get('name'))[:160]
                if not re.fullmatch(r'[a-zA-Z0-9_-]{1,80}', sid) or not name or sid in seen or norm(name) in names:
                    raise ValueError('Fornecedores devem ter nomes e identificações únicos.')
                phone = re.sub(r'[^\d+]', '', text(s.get('phone')))
                if phone and not re.fullmatch(r'\+?[1-9]\d{7,14}', phone):
                    raise ValueError('Informe WhatsApp com país e DDD, por exemplo 5585999999999.')
                suppliers.append({'id': sid, 'name': name, 'phone': phone})
                seen.add(sid)
                names.add(norm(name))
            data['suppliers'] = suppliers

            notes = {i['id']: i for i in body.get('items', [])}
            for i in data['items']:
                if i['id'] in notes:
                    i['note'] = text(notes[i['id']].get('note'))[:2000]
                    i['purchase_type'] = text(notes[i['id']].get('purchase_type'))[:100]

            raw_quotes = body.get('quotes', {})
            data['quotes'] = {i['id']: {s['id']: raw_quotes.get(i['id'], {}).get(s['id'], {}) for s in suppliers} for i in data['items']}
            for quotes in data['quotes'].values():
                for sid, q in quotes.items():
                    normalized = {'price': '' if q.get('price') in ('', None) else str(number(q['price']))}
                    delivery = text(q.get('delivery'))[:120]
                    if delivery:
                        normalized['delivery'] = delivery
                    if 'negotiated' in q:
                        # Preço inicial vazio significa que este fornecedor não cotou este item.
                        # Qualquer valor negociado residual é descartado para não participar da comparação.
                        normalized['negotiated'] = '' if not normalized['price'] or q['negotiated'] in ('', None) else str(number(q['negotiated']))
                    else:
                        normalized.update(discount=str(number(q.get('discount') or 0)), kind=q.get('kind', 'percent'))
                    quotes[sid] = normalized

            # Existing `choices` remains the compatibility mechanism for old tie-only maps.
            data['choices'] = body.get('choices', data.get('choices', {})) or {}
            if removed or added:
                data['choices'] = {iid: sid for iid, sid in data['choices'].items() if iid not in removed | added}

            valid_items = {i['id'] for i in data['items']}
            valid_suppliers = {s['id'] for s in suppliers}
            raw_awards = body.get('awards', data.get('awards', {})) or {}
            awards = {}
            if not isinstance(raw_awards, dict):
                raise ValueError('Escolha de fornecedor inválida.')
            for item_id, award in raw_awards.items():
                if item_id not in valid_items or not isinstance(award, dict):
                    continue
                supplier_id = text(award.get('supplier_id'))
                if not supplier_id:
                    continue
                if supplier_id not in valid_suppliers:
                    raise ValueError('Fornecedor escolhido não pertence mais a este mapa.')
                quote = data['quotes'].get(item_id, {}).get(supplier_id, {})
                if quote.get('price') in ('', None):
                    raise ValueError('Só é possível escolher um fornecedor que tenha cotado o item.')
                reason = text(award.get('reason'))[:120]
                note = text(award.get('note'))[:500]
                if reason and reason not in AWARD_REASONS:
                    raise ValueError('Motivo da escolha inválido.')
                awards[item_id] = {'supplier_id': supplier_id, 'reason': reason, 'note': note}
            data['awards'] = awards

            # Existing rows retain their quotes/decisions. Added rows always start without prices.
            data['items'].extend(new_items)
            for item in new_items:
                data['quotes'][item['id']] = {s['id']: {'price': '', 'negotiated': ''} for s in suppliers}
            result = evaluate(data)
            for line in result['lines']:
                if line['manual_selection'] and Decimal(line['opportunity_cost']) > 0 and not line['selection_reason']:
                    raise ValueError('Informe o motivo ao escolher uma proposta acima do menor preço.')

            data['revision'] += 1
            data['updated'] = now()
            if removed or added:
                self.create_backup(reason='pre-remove-map-items' if removed else 'pre-add-map-items', automatic=True)
            self.put('maps', data['id'], data)
        return self.detail(data['id'])

    def detail(self, mid):
        data = self.get_map(mid)
        # Stable hotel order also applies to existing maps. Prices/decisions remain keyed by item ID.
        data['items'] = sorted(data['items'], key=lambda item: norm(item.get('company')))
        current = {i['id']: i for i in self.all('items')}
        changes = []
        for item in data['items']:
            fresh = current.get(item['id'])
            if not fresh:
                changes.append(f"SCI {item['sci']} • {item['description']}: não consta mais na lista de itens elegíveis.")
            elif any(item[k] != fresh[k] for k in ('quantity', 'unit', 'description', 'company', 'buyer', 'approval')):
                changes.append(f"SCI {item['sci']} • {item['description']}: dados alterados na última importação.")
        return {'map': data, 'result': evaluate(data), 'changes': changes}

    def preview(self, mid, sid):
        detail = self.detail(mid)
        data = detail['map']
        if data['archived']:
            raise ValueError('Mapa concluído. Crie um novo mapa para solicitar cotações.')
        if detail['changes']:
            raise ValueError('A base mudou para itens deste mapa. Conclua este mapa e crie outro com os dados atuais antes de enviar.')
        supplier = next((s for s in data['suppliers'] if s['id'] == sid), None)
        if not supplier:
            raise ValueError('Fornecedor não encontrado.')
        lines = [f"Olá, {supplier['name']}! Poderia cotar os itens abaixo?", f"Mapa: {data['name']}", '']
        previous_hotel = None
        for n, item in enumerate(data['items'], 1):
            cnpj = hotel_cnpj(item['company'])
            if not cnpj:
                raise ValueError(f"CNPJ não cadastrado para o hotel {item['company']}. Atualize a associação antes de enviar a cotação.")
            hotel = norm(item['company'])
            if hotel != previous_hotel:
                lines += [f"Hotel: {item['company']} | CNPJ: {cnpj}", '']
                previous_hotel = hotel
            lines += [f"{n}. {item['description']}", f"Quantidade: {item['quantity']} {item['unit']}"]
            if item['note']:
                lines.append('Observação: ' + item['note'])
            lines.append('')
        lines.append('Por favor, informe preço unitário, disponibilidade, prazo de entrega, frete, condições de pagamento e eventual desconto. Obrigado!')
        message = '\n'.join(lines)
        if len(message) > 60000:
            raise ValueError('Mapa muito grande para uma mensagem. Divida os itens em mapas menores.')
        digest = hashlib.sha256((mid + '|' + supplier['phone'] + '|' + message).encode()).hexdigest()
        return {'supplier': supplier, 'message': message, 'fingerprint': digest, 'revision': data['revision']}

    def negotiation_preview(self, mid, sid):
        detail = self.detail(mid)
        data = detail['map']
        result = detail['result']
        if data['archived']:
            raise ValueError('Mapa concluído. Negociações só podem ser enviadas em mapas em cotação.')
        if detail['changes']:
            raise ValueError('A base mudou para itens deste mapa. Conclua este mapa e crie outro com os dados atuais antes de negociar.')
        supplier = next((s for s in data['suppliers'] if s['id'] == sid), None)
        if not supplier:
            raise ValueError('Fornecedor não encontrado.')

        targets = []
        items_by_id = {i['id']: i for i in data['items']}
        for line in result['lines']:
            quote = next((q for q in line['quotes'] if q['supplier_id'] == sid), None)
            if not quote or not line.get('target_price'):
                continue
            current = Decimal(quote['final_price'])
            target = Decimal(line['target_price'])
            if current <= target:
                continue
            item = items_by_id[line['id']]
            targets.append({
                'item_id': item['id'],
                'description': item['description'],
                'quantity': item['quantity'],
                'unit': item['unit'],
                'current_price': str(current),
                'target_price': str(target),
                'reduction_unit': quote['target_reduction_unit'],
                'required_reduction_percent': quote['required_reduction_percent'],
            })

        if not targets:
            raise ValueError('Este fornecedor não possui itens pendentes de negociação para a meta atual.')

        singular = len(targets) == 1
        if singular:
            lines = ['Para fecharmos, preciso que você baixe um pouco esse valor. Consegue fechar esse item no valor abaixo?', '']
        else:
            lines = ['Para fecharmos, preciso que você baixe um pouco esses valores. Consegue fechar os itens nos valores abaixo?', '']
        for target in targets:
            lines.append(f"{target['description']} — {target['quantity']} {target['unit']}")
            lines.append(f"{format_brl(target['target_price'])} unitário")
            lines.append('')
        lines.append('Se conseguir chegar nesse valor, me confirma por favor.' if singular else 'Se conseguir chegar nesses valores, me confirma por favor.')
        message = '\n'.join(lines)
        if len(message) > 60000:
            raise ValueError('Negociação muito grande para uma mensagem. Divida os itens em mapas menores.')
        # Fingerprint intentionally binds the preview to the current map revision and target values.
        context = json.dumps(targets, ensure_ascii=False, sort_keys=True)
        digest = hashlib.sha256((mid + '|' + sid + '|negotiation|' + str(data['revision']) + '|' + context).encode()).hexdigest()
        return {'supplier': supplier, 'message': message, 'fingerprint': digest,
                'revision': data['revision'], 'saving_target': result['saving_target_percent'], 'targets': targets}

    def send_negotiation(self, body):
        if not self.send_lock.acquire(blocking=False):
            raise ValueError('Já existe um envio em andamento.')
        record = None
        submitted = False
        try:
            with self.lock:
                p = self.negotiation_preview(body['map_id'], body['supplier_id'])
                if p['fingerprint'] != body.get('fingerprint') or p['revision'] != body.get('revision'):
                    raise ValueError('A meta ou os valores do mapa mudaram. Atualize a prévia antes de enviar.')
                if not p['supplier']['phone']:
                    raise ValueError('Cadastre o WhatsApp deste fornecedor.')
                message = text(body.get('message'))
                if not message:
                    raise ValueError('A mensagem de negociação não pode ficar vazia.')
                if len(message) > 60000:
                    raise ValueError('Mensagem muito grande para o WhatsApp.')
                send_fingerprint = hashlib.sha256((body['map_id'] + '|' + p['supplier']['phone'] + '|negotiation|' + message).encode()).hexdigest()
                if any(m.get('fingerprint') == send_fingerprint and m.get('status') in ('sent', 'sending', 'uncertain') for m in self.all('messages')):
                    raise ValueError('Esta negociação já foi enviada ou está incerta. Consulte o Histórico.')
                record = {'id': uuid.uuid4().hex, 'at': now(), 'map_id': body['map_id'], 'kind': 'negotiation',
                          'supplier': p['supplier'], 'message': message, 'fingerprint': send_fingerprint,
                          'revision': p['revision'], 'targets': p['targets'], 'saving_target': p['saving_target'],
                          'status': 'sending'}
                self.put('messages', record['id'], record)
            whatsapp_request('/wait')
            submitted = True
            response = whatsapp_request('/send', {'phone': p['supplier']['phone'], 'message': message})
            record.update(status=response.get('status', 'uncertain'), error=response.get('error', ''), message_id=response.get('message_id'))
            if record['status'] not in ('sent', 'failed', 'uncertain'):
                record['status'] = 'uncertain'
        except Exception as exc:
            if record is None:
                raise
            record.update(status='uncertain' if submitted else 'failed', error=str(exc))
        finally:
            if record:
                self.put('messages', record['id'], record)
            self.send_lock.release()
        return record

    def send(self, body):
        if not self.send_lock.acquire(blocking=False):
            raise ValueError('Já existe um envio em andamento.')
        record = None
        submitted = False
        try:
            with self.lock:
                p = self.preview(body['map_id'], body['supplier_id'])
                if p['fingerprint'] != body.get('fingerprint') or p['revision'] != body.get('revision'):
                    raise ValueError('Mensagem alterada. Atualize a prévia antes de enviar.')
                if not p['supplier']['phone']:
                    raise ValueError('Cadastre o WhatsApp deste fornecedor.')
                if any(m['fingerprint'] == p['fingerprint'] and m['status'] in ('sent', 'sending', 'uncertain') for m in self.all('messages')):
                    raise ValueError('Esta solicitação já foi enviada ou está incerta. Consulte o Histórico.')
                record = {'id': uuid.uuid4().hex, 'at': now(), 'map_id': body['map_id'], 'kind': 'quote', **p, 'status': 'sending'}
                self.put('messages', record['id'], record)
            whatsapp_request('/wait')
            submitted = True
            response = whatsapp_request('/send', {'phone': p['supplier']['phone'], 'message': p['message']})
            record.update(status=response.get('status', 'uncertain'), error=response.get('error', ''), message_id=response.get('message_id'))
            if record['status'] not in ('sent', 'failed', 'uncertain'):
                record['status'] = 'uncertain'
        except Exception as exc:
            if record is None:
                raise
            record.update(status='uncertain' if submitted else 'failed', error=str(exc))
        finally:
            if record:
                self.put('messages', record['id'], record)
            self.send_lock.release()
        return record

    def review(self, body):
        with self.lock:
            msg = next((m for m in self.all('messages') if m['id'] == body['id']), None)
            if not msg or msg['status'] != 'uncertain' or body.get('outcome') not in ('received', 'not_received'):
                raise ValueError('Revisão inválida.')
            msg.update(status='sent' if body['outcome'] == 'received' else 'reviewed_not_received', reviewed_at=now())
            self.put('messages', msg['id'], msg)
        return msg

    def export_csv(self, mid):
        data = self.detail(mid)
        m = data['map']
        rows = [['HOTEL', 'SCI', 'TIPO DE COMPRA', 'DESCRIÇÃO', 'QTD', 'UNIDADE', 'OBSERVAÇÃO'] +
                [f"{s['name']} — {label}" for s in m['suppliers'] for label in ('Preço unitário', 'Total bruto', 'Desconto R$', 'Total líquido')] +
                ['FORNECEDOR ESCOLHIDO', 'VALOR FINAL', 'ECONOMIA R$', 'PRAZO ESCOLHIDO', 'MOTIVO DA ESCOLHA', 'OBSERVAÇÃO DA ESCOLHA', 'DIF. VS MENOR PREÇO R$']]
        for i, result in zip(m['items'], data['result']['lines']):
            row = [i['company'], i['sci'], i['purchase_type'], i['description'], i['quantity'], i['unit'], i['note']]
            for s in m['suppliers']:
                q = next((q for q in result['quotes'] if q['supplier_id'] == s['id']), None)
                row += [str(m['quotes'][i['id']][s['id']]['price']).replace('.', ','), q['gross'].replace('.', ','), q['saving'].replace('.', ','), q['net'].replace('.', ',')] if q else ['', '', '', '']
            chosen = result['chosen']
            if chosen:
                row += [chosen['supplier'], chosen['net'].replace('.', ','), chosen['saving'].replace('.', ','),
                        chosen.get('delivery', ''), result.get('selection_reason', ''), result.get('selection_note', ''),
                        result.get('opportunity_cost', '0').replace('.', ',')]
            else:
                row += [('EMPATE' if result['state'] == 'tie' else 'SEM COTAÇÃO'), '', '', '', '', '', '']
            rows.append(row)
        out = io.StringIO()
        writer = csv.writer(out, delimiter=';')
        for row in rows:
            writer.writerow(["'" + str(v) if str(v).lstrip().startswith(('=', '+', '-', '@')) else v for v in row])
        return {'filename': 'mapa-' + mid[:8] + '.csv', 'content': '\ufeff' + out.getvalue()}

    def export_xls(self, mid):
        from excel_export import build_xls
        detail = self.detail(mid)
        raw = build_xls(detail['map'], detail['result'])
        return {'filename': 'mapa-' + mid[:8] + '.xls', 'encoding': 'base64',
                'content': base64.b64encode(raw).decode('ascii')}


class Handler(BaseHTTPRequestHandler):
    def log_message(self, *args):
        pass

    def reply(self, status, data):
        raw = json.dumps(data, ensure_ascii=False).encode()
        self.send_response(status)
        self.send_header('Content-Type', 'application/json; charset=utf-8')
        self.send_header('Content-Length', str(len(raw)))
        self.send_header('Cache-Control', 'no-store')
        self.end_headers()
        self.wfile.write(raw)

    def run_request(self):
        token = os.environ.get('FOLLOWUP_API_TOKEN', '')
        if not token or not hmac.compare_digest(self.headers.get('X-FollowUp-Token', ''), token):
            return self.reply(401, {'error': 'Não autorizado.'})
        try:
            parsed = urlparse(self.path)
            query = {k: v[0] for k, v in parse_qs(parsed.query).items()}
            s = self.server.store
            if self.command == 'GET':
                if parsed.path == '/health':
                    return self.reply(200, {'ok': True})
                if parsed.path == '/overview':
                    return self.reply(200, s.overview())
                if parsed.path == '/items':
                    return self.reply(200, s.catalog())
                if parsed.path == '/maps':
                    return self.reply(200, s.map_summaries())
                if parsed.path == '/map':
                    return self.reply(200, s.detail(query['id']))
                if parsed.path == '/settings':
                    return self.reply(200, s.settings())
                if parsed.path == '/data-safety':
                    return self.reply(200, s.data_safety_status())
                if parsed.path == '/preview':
                    return self.reply(200, s.preview(query['id'], query['supplier']))
                if parsed.path == '/negotiation-preview':
                    return self.reply(200, s.negotiation_preview(query['id'], query['supplier']))
                if parsed.path == '/history':
                    return self.reply(200, s.all('messages')[::-1])
                if parsed.path == '/export':
                    return self.reply(200, s.export_xls(query['id']))
            else:
                size = int(self.headers.get('Content-Length', '0'))
                if size > 5000000:
                    raise ValueError('Requisição muito grande.')
                body = json.loads(self.rfile.read(size) or '{}')
                if parsed.path == '/import':
                    return self.reply(200, s.import_file(body['path']))
                if parsed.path == '/maps/create':
                    return self.reply(200, s.create_map(body))
                if parsed.path == '/maps/save':
                    return self.reply(200, s.save_map(body))
                if parsed.path in ('/maps/complete', '/maps/archive'):
                    # /maps/archive remains as a compatibility alias for older renderers.
                    return self.reply(200, s.complete_map(body['id']))
                if parsed.path == '/maps/delete':
                    return self.reply(200, s.delete_map(body['id']))
                if parsed.path == '/settings':
                    if not isinstance(body, dict):
                        raise ValueError('Configuração inválida.')
                    body['id'] = 'ui'
                    s.put('settings', 'ui', body)
                    return self.reply(200, body)
                if parsed.path == '/data-safety/backup':
                    return self.reply(200, s.create_backup(
                        reason=body.get('reason', 'manual'), automatic=bool(body.get('automatic', False))
                    ))
                if parsed.path == '/send':
                    return self.reply(200, s.send(body))
                if parsed.path == '/send-negotiation':
                    return self.reply(200, s.send_negotiation(body))
                if parsed.path == '/review':
                    return self.reply(200, s.review(body))
            self.reply(404, {'error': 'Operação não encontrada.'})
        except Exception as exc:
            self.reply(400, {'error': str(exc)})

    do_GET = run_request
    do_POST = run_request


if __name__ == '__main__':
    parser = argparse.ArgumentParser()
    parser.add_argument('command', choices=['serve', 'summary'])
    parser.add_argument('--port', type=int, default=0)
    args = parser.parse_args()
    try:
        store = Store(os.environ.get('FOLLOWUP_DATA_DIR', str(Path.home() / '.vyzium-compras')))
    except DataIntegrityError as exc:
        print(f'VYZIUM_DATA_INTEGRITY: {exc}', file=sys.stderr, flush=True)
        raise SystemExit(3)
    if args.command == 'summary':
        print(json.dumps(store.overview(), ensure_ascii=False), flush=True)
    else:
        server = ThreadingHTTPServer(('127.0.0.1', args.port), Handler)
        server.store = store
        print(json.dumps({'event': 'ready', 'port': server.server_address[1]}), flush=True)
        server.serve_forever()
